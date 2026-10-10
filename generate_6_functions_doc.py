import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_document():
    doc = docx.Document()

    # Page Margins (0.8 inch top/bottom, 0.85 inch left/right)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    COLOR_PRIMARY = RGBColor(31, 78, 120)     # #1F4E78 Navy
    COLOR_SECONDARY = RGBColor(46, 117, 182)  # #2E75B6 Blue
    COLOR_ACCENT = RGBColor(16, 185, 129)     # #10B981 Emerald
    COLOR_TEXT = RGBColor(30, 41, 59)         # #1E293B Dark Slate
    COLOR_MUTED = RGBColor(100, 116, 139)     # #64748B Muted

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(3)

    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        cell._tc.get_or_add_tcPr().append(tcMar)

    def set_table_borders(table, color="CBD5E1"):
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            </w:tblBorders>
        ''')
        table._tbl.tblPr.append(borders)

    def add_callout(text, title=None, alert_type="note"):
        bg_color = "EFF6FF"
        border_color = "2563EB"
        icon = "📌"
        if alert_type == "success":
            bg_color = "ECFDF5"
            border_color = "10B981"
            icon = "✅"
        elif alert_type == "warning":
            bg_color = "FFFBEB"
            border_color = "F59E0B"
            icon = "⚠️"

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=130, bottom=130, left=180, right=180)

        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
                <w:top w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        if title:
            r_title = p.add_run(f"{icon} {title}\n")
            r_title.bold = True
            r_title.font.name = "Segoe UI"
            r_title.font.size = Pt(10.5)
            r_title.font.color.rgb = RGBColor(30, 58, 138) if alert_type == "note" else RGBColor(6, 95, 70)
        r_text = p.add_run(text)
        r_text.font.name = "Segoe UI"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = COLOR_TEXT

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(2)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Segoe UI'
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10)
        r.font.italic = italic
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Segoe UI'
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_TEXT
        return p

    def render_function_table(sub_functions):
        table = doc.add_table(rows=len(sub_functions) + 1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)

        col_widths = [1.2, 1.4, 1.3, 1.6, 1.3]
        headers = ["Chức năng nhỏ", "Mục đích nghiệp vụ", "Dữ liệu đầu vào (Input)", "Xử lý & Ràng buộc (Processing)", "Kết quả đầu ra (Output)"]

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].width = Inches(col_widths[i])
            set_cell_background(hdr_cells[i], "1F4E78")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
            p = hdr_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h)
            r.bold = True
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, f in enumerate(sub_functions):
            row_cells = table.rows[r_idx + 1].cells
            bg_col = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx, val in enumerate(f):
                row_cells[c_idx].width = Inches(col_widths[c_idx])
                set_cell_background(row_cells[c_idx], bg_col)
                set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(val))
                r.font.name = 'Segoe UI'
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(4)

    # =========================================================================
    # BANNER & TIÊU ĐỀ
    # =========================================================================
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_pre.add_run("BÁO CÁO CHI TIẾT ĐẶC TẢ CHỨC NĂNG PHẦN MỀM - BÀI TẬP LỚN LẬP TRÌNH CSDL\n")
    r_uni.font.name = "Segoe UI"
    r_uni.font.size = Pt(11)
    r_uni.bold = True
    r_uni.font.color.rgb = COLOR_SECONDARY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("ĐẶC TẢ CHI TIẾT 6 CHỨC NĂNG NHỎ TRONG 12 FORM LỚN\nHỆ THỐNG QUẢN LÝ KHU VUI CHƠI GIẢI TRÍ (WONDERLAND PARK / FUNZONE)")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Phân rã chi tiết 72 chức năng con | Input - Thuật toán - Validation - Output | 4 Thành viên")
    r_sub.italic = True
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = COLOR_MUTED

    add_callout(
        "Tài liệu này đặc tả chi tiết 6 chức năng nhỏ (Sub-functions / Core Operations) bên trong từng Form thuộc 12 Forms lớn bắt buộc của dự án Wonderland Park.\n"
        "Mỗi chức năng con đều được phân tích theo 5 tiêu chí kỹ thuật: (1) Tên chức năng con, (2) Mục đích nghiệp vụ, (3) Dữ liệu đầu vào (Input/Trigger), (4) Thuật toán xử lý & Kiểm tra ràng buộc (Processing/Validation/Transaction), (5) Kết quả đầu ra (Output/Impact CSDL).\n"
        "Tổng cộng tài liệu phân rã chi tiết 72 chức năng con bao phủ toàn diện 100% mã nguồn thực tế.",
        title="TỔNG QUAN TÀI LIỆU PHÂN RÃ CHỨC NĂNG",
        alert_type="note"
    )

    doc.add_page_break()

    # =========================================================================
    # DANH MỤC 12 FORMS THEO 4 THÀNH VIÊN
    # =========================================================================
    add_h1("TỔNG HỢP DANH MỤC 12 FORM LỚN THEO 4 THÀNH VIÊN")
    add_p("Hệ thống phân chia đồng đều 12 Form cho 4 thành viên, mỗi thành viên đảm nhiệm trọn vẹn 3 Form lớn:")

    overview_table = doc.add_table(rows=5, cols=4)
    overview_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    overview_table.autofit = False
    set_table_borders(overview_table)
    col_w = [1.3, 1.7, 1.8, 2.0]
    headers_ov = ["Thành viên", "Phân hệ đảm nhiệm", "Danh sách 3 Form lớn", "Quy mô chức năng con"]
    hdr_cells = overview_table.rows[0].cells
    for i, h in enumerate(headers_ov):
        hdr_cells[i].width = Inches(col_w[i])
        set_cell_background(hdr_cells[i], "1F4E78")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.name = 'Segoe UI'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_ov = [
        ["THÀNH VIÊN 1", "Quản lý Khách hàng", "• Form 1: FormKhachHang\n• Form 2: FormImportKhachHang\n• Form 3: FormExportKhachHang", "18 chức năng con (CRUD, Tra cứu đa chế độ, Pipeline Import kiểm lỗi từng dòng, Export ClosedXML)."],
        ["THÀNH VIÊN 2", "Quản lý Trò chơi & Dịch vụ", "• Form 4: FormTroChoi\n• Form 5: FormDichVu\n• Form 6: FormImportExportTroChoiDichVu", "18 chức năng con (Quy chuẩn an toàn, Sức chứa, Quản lý kho hàng, Cập nhật tồn nhanh, TabControl Import/Export)."],
        ["THÀNH VIÊN 3", "Quản lý Nhân viên & Bảng giá Vé", "• Form 7: FormNhanVien\n• Form 8: FormVe\n• Form 9: FormImportExportNhanVienVe", "18 chức năng con (Hồ sơ nhân sự, Phân quyền, Bảng giá vé ngày/VIP, TabControl nạp/xuất dữ liệu mẫu)."],
        ["THÀNH VIÊN 4", "Bán vé POS, Hóa đơn & Báo cáo", "• Form 10: FormBanVe (POS)\n• Form 11: FormHoaDon\n• Form 12: FormBaoCaoDoanhThu", "18 chức năng con (POS Card, Giỏ hàng động, Chiết khấu, Trừ tồn kho trong Transaction, Hủy đơn hoàn kho, Báo cáo tài chính)."]
    ]
    for r_idx, row in enumerate(data_ov):
        row_cells = overview_table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].width = Inches(col_w[c_idx])
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

    # =========================================================================
    # PHẦN 1: THÀNH VIÊN 1 - QUẢN LÝ KHÁCH HÀNG
    # =========================================================================
    add_h1("PHẦN 1: CHI TIẾT 6 CHỨC NĂNG NHỎ TRONG CÁC FORM CỦA THÀNH VIÊN 1")

    # Form 1: FormKhachHang
    add_h2("1.1. FORM 1: FormKhachHang (Quản lý Hồ sơ Khách hàng)")
    add_p("Form đảm nhiệm quản lý hồ sơ thông tin khách hàng thân thiết, lịch sử đăng ký thành viên và hỗ trợ tra cứu đa phương thức.")
    
    f1_subs = [
        [
            "1. Thêm mới Khách hàng\n(Create / Add)",
            "Tiếp nhận và lưu trữ hồ sơ khách hàng mới đến vui chơi.",
            "• Họ tên, Ngày sinh, Giới tính, Số điện thoại, Email, Địa chỉ.",
            "• Tự sinh mã kế tiếp thông minh dạng KHxxx (ví dụ: KH021).\n• Validation: Bắt buộc Họ tên và SĐT; SĐT đủ 10 số (0xxxxxxxxx); Email chuẩn regex; Kiểm tra trùng SĐT trong CSDL.\n• Gọi CustomerService.Add(khachHang).",
            "• Thêm bản ghi vào bảng KhachHang.\n• Hiển thị MessageBox thông báo thành công.\n• Tự nạp lại dữ liệu lên DataGridView."
        ],
        [
            "2. Đồng bộ khi chọn dòng\n(CellClick / Data Binding)",
            "Tự động đưa thông tin dòng đang chọn lên form nhập liệu.",
            "• Sự kiện người dùng click vào dòng trên DataGridView.",
            "• Lấy MaKH từ dòng được chọn.\n• Khóa ô txtMaKH (ReadOnly = true) để chống sửa khóa chính.\n• Điền Họ tên, Ngày sinh, Giới tính, SĐT, Email, Địa chỉ lên controls.",
            "• Người dùng quan sát toàn bộ thông tin chi tiết của khách hàng sẵn sàng để sửa hoặc tra cứu."
        ],
        [
            "3. Cập nhật thông tin\n(Update Customer)",
            "Chỉnh sửa thông tin liên lạc, địa chỉ của khách hàng.",
            "• Dữ liệu đã thay đổi trên các ô nhập liệu.",
            "• Kiểm tra Họ tên và SĐT không rỗng.\n• Kiểm tra định dạng SĐT (10 số) và Email.\n• Kiểm tra trùng SĐT: SĐT không được trùng với khách hàng khác trong hệ thống.\n• Cập nhật bản ghi qua CustomerService.Update().",
            "• Cập nhật thông tin trong SQL Server.\n• Thông báo 'Cập nhật thành công!'.\n• Làm mới bảng hiển thị."
        ],
        [
            "4. Xóa Khách hàng an toàn\n(Safe Delete)",
            "Xóa hồ sơ khách hàng nhưng phải đảm bảo toàn vẹn dữ liệu.",
            "• Nhấn nút [Xóa], MaKH của khách hàng đang chọn.",
            "• Hộp thoại xác nhận MessageBox Yes/No.\n• Kiểm tra toàn vẹn tham chiếu (FK Check): Truy vấn bảng HoaDon. Nếu khách hàng đã từng có hóa đơn mua hàng (existing.HoaDons.Any()) -> Chặn lại và cảnh báo 'Không thể xóa vì tồn tại N hóa đơn liên quan!'.\n• Nếu chưa có hóa đơn: _context.KhachHangs.Remove().",
            "• Khách hàng được xóa an toàn khỏi CSDL nếu chưa phát sinh giao dịch.\n• Đảm bảo tuyệt đối không lỗi Crash và không vi phạm toàn vẹn dữ liệu."
        ],
        [
            "5. Tìm kiếm đa tiêu chí\n(Multi-mode Search)",
            "Tra cứu nhanh khách hàng theo từ khóa và chế độ tìm kiếm.",
            "• Từ khóa tìm kiếm (txtSearch), Kiểu tìm (Contains / StartsWith / Exact).",
            "• LINQ truy vấn đồng thời trên 4 trường: MaKH, HoTen, SoDienThoai, Email.\n• So khớp chuỗi theo tùy chọn: Chứa (Contains), Bắt đầu bằng (StartsWith), Chính xác (Exact Match).",
            "• DataGridView chỉ hiển thị các khách hàng thỏa mãn điều kiện.\n• Cập nhật Label tổng số lượng bản ghi tìm thấy."
        ],
        [
            "6. Làm mới / Xóa trắng form\n(Reset / Clear Form)",
            "Đưa form nhập liệu về trạng thái ban đầu để nhập mới.",
            "• Nhấn nút [Làm mới].",
            "• Xóa trắng các TextBox: HoTen, SoDienThoai, Email, DiaChi.\n• Đặt lại Ngày sinh về mặc định.\n• Mở khóa và tự sinh mã mới kế tiếp vào txtMaKH.\n• Tải lại toàn bộ danh sách khách hàng mới nhất từ CSDL.",
            "• Giao diện sạch sẽ, sẵn sàng cho việc nhập khách hàng tiếp theo."
        ]
    ]
    render_function_table(f1_subs)

    # Form 2: FormImportKhachHang
    add_h2("1.2. FORM 2: FormImportKhachHang (Nhập Khách hàng từ Excel)")
    add_p("Form xử lý quy trình nạp hàng loạt khách hàng từ file Excel .xlsx bằng ClosedXML với pipeline kiểm tra lỗi từng dòng nghiêm ngặt.")

    f2_subs = [
        [
            "1. Tải file mẫu Excel\n(Download Template)",
            "Cung cấp file Excel chuẩn cấu trúc cột cho người dùng nhập liệu.",
            "• Nhấn nút [Tải file mẫu Excel].",
            "• Sử dụng TemplateGenerator sinh file KhachHang_Template.xlsx.\n• Định dạng sẵn các cột: HoTen, NgaySinh, GioiTinh, SoDienThoai, Email, DiaChi kèm dữ liệu mẫu.\n• Mở SaveFileDialog cho người dùng lưu file.",
            "• File template chuẩn cấu trúc được lưu về máy tính của người dùng."
        ],
        [
            "2. Chọn & Duyệt file Excel\n(Browse .xlsx File)",
            "Tiếp nhận file dữ liệu Excel người dùng muốn nạp vào hệ thống.",
            "• Nhấn nút [Chọn file Excel], chọn file .xlsx từ OpenFileDialog.",
            "• Kiểm tra đuôi file (*.xlsx).\n• Kiểm tra sự tồn tại và tính hợp lệ của file trên ổ đĩa.\n• Hiển thị đường dẫn đầy đủ lên txtFilePath.",
            "• Đường dẫn file được nạp vào bộ nhớ sẵn sàng cho quá trình đọc dữ liệu."
        ],
        [
            "3. Đọc dữ liệu ClosedXML\n(Excel Parsing Engine)",
            "Mở và bóc tách dữ liệu từ file Excel vào bộ nhớ ứng dụng.",
            "• Sự kiện sau khi chọn file.",
            "• Sử dụng ClosedXML (XLWorkbook) mở file.\n• Truy xuất Worksheet đầu tiên.\n• Duyệt từ dòng thứ 2 (bỏ qua dòng tiêu đề) đến LastRowUsed().\n• Đọc giá trị từng ô và ép kiểu dữ liệu an toàn (DateTime, String).",
            "• Tạo danh sách thô các đối tượng khách hàng (List<KhachHang>)."
        ],
        [
            "4. Kiểm tra lỗi từng dòng\n(Row Validation Engine)",
            "Kiểm tra tính hợp lệ dữ liệu và bắt lỗi chi tiết từng hàng.",
            "• Tập dữ liệu thô vừa đọc từ file Excel.",
            "• Duyệt từng dòng: Kiểm tra Họ tên và SĐT không rỗng; SĐT đúng 10 số; Email đúng định dạng; Kiểm tra không trùng mã hoặc trùng SĐT đã có trong SQL Server; Kiểm tra không trùng lặp SĐT ngay trong file Excel.\n• Nếu có lỗi: Ghi lại lỗi chi tiết ('Dòng X: Số điện thoại không đúng 10 số').",
            "• Phân loại thành: Danh sách dòng hợp lệ (ValidItems) và Danh sách thông báo lỗi (ErrorList)."
        ],
        [
            "5. Xem trước (Preview) & Báo lỗi\n(Preview & Error List)",
            "Cho người dùng kiểm tra trực quan dữ liệu trước khi quyết định lưu.",
            "• Kết quả từ bước kiểm tra lỗi.",
            "• Đổ danh sách dòng hợp lệ lên DataGridView Preview.\n• Đổ danh sách lỗi chi tiết lên ListBox lstErrors.\n• Cập nhật nhãn tổng kết: 'Hợp lệ: X dòng | Lỗi: Y dòng'.\n• Kích hoạt nút [Xác nhận nạp] nếu có ít nhất 1 dòng hợp lệ.",
            "• Người dùng nhìn rõ dữ liệu nào sẽ vào CSDL, dữ liệu nào bị từ chối."
        ],
        [
            "6. Xác nhận nạp vào CSDL\n(Confirm Bulk Insert)",
            "Ghi toàn bộ các dòng hợp lệ vào cơ sở dữ liệu SQL Server.",
            "• Nhấn nút [Xác nhận nạp vào CSDL].",
            "• Duyệt qua từng dòng trong ValidItems: Tự động sinh mã KHxxx kế tiếp duy nhất.\n• Gọi CustomerService.AddBatch() hoặc AddRange() trong Transaction.\n• Gọi SaveChanges() ghi vào bảng KhachHang.",
            "• Lưu thành công toàn bộ khách hàng hợp lệ vào SQL Server.\n• Thông báo số lượng bản ghi thêm mới thành công."
        ]
    ]
    render_function_table(f2_subs)

    # Form 3: FormExportKhachHang
    add_h2("1.3. FORM 3: FormExportKhachHang (Xuất Khách hàng ra Excel)")
    add_p("Form phụ trách xuất danh sách khách hàng ra file Excel định dạng báo cáo doanh nghiệp chuẩn mực bằng thư viện ClosedXML.")

    f3_subs = [
        [
            "1. Lựa chọn phạm vi xuất\n(Select Export Scope)",
            "Cho phép người dùng chọn xuất toàn bộ hoặc chỉ xuất kết quả lọc.",
            "• Chọn RadioButton: 'Xuất toàn bộ' hoặc 'Xuất theo bộ lọc hiện tại'.",
            "• Nếu chọn toàn bộ: Gọi CustomerService.GetAll().\n• Nếu chọn theo bộ lọc: Lấy danh sách đang hiển thị theo từ khóa tìm kiếm.\n• Kiểm tra nếu danh sách rỗng -> Thông báo không có dữ liệu để xuất.",
            "• Tập dữ liệu khách hàng được nạp vào sẵn sàng xuất file."
        ],
        [
            "2. Tạo Banner Header báo cáo\n(Report Header Styling)",
            "Thiết kế phần đầu file Excel chuyên nghiệp, sang trọng.",
            "• Khởi tạo XLWorkbook và tạo Worksheet 'Khách Hàng'.",
            "• Merge ô từ A1 đến H1 làm Banner tiêu đề.\n• Đặt nền màu xanh Navy #1F4E78, chữ in hoa màu trắng, cỡ chữ 16pt, in đậm, căn giữa.\n• Dòng A2: Ghi thời gian xuất file ('Thời gian xuất: dd/MM/yyyy HH:mm') và tổng số bản ghi.",
            "• Tiêu đề báo cáo chuẩn doanh nghiệp ở đầu trang tính Excel."
        ],
        [
            "3. Thiết lập tiêu đề cột\n(Column Header Formatting)",
            "Tạo hàng tiêu đề cho bảng dữ liệu với phong cách nổi bật.",
            "• Danh sách 8 cột: STT, Mã KH, Họ tên, Ngày sinh, Giới tính, Số điện thoại, Email, Địa chỉ.",
            "• Thiết lập nền màu xanh biển #2E75B6 cho dòng tiêu đề cột.\n• Chữ in hoa màu trắng, in đậm, cỡ chữ 11pt.\n• Kẻ viền khung mỏng (Thin border) xung quanh các ô tiêu đề.",
            "• Hàng tiêu đề cột rõ ràng, bắt mắt và dễ đọc."
        ],
        [
            "4. Định dạng dữ liệu từng ô\n(Cell Data Formatting)",
            "Điền dữ liệu và định dạng chuẩn theo từng loại dữ liệu.",
            "• Dữ liệu từng dòng khách hàng.",
            "• Cột Ngày sinh: Định dạng hiển thị 'dd/MM/yyyy'.\n• Cột Giới tính, Mã KH, SĐT: Định dạng căn giữa.\n• Cột Họ tên, Email, Địa chỉ: Định dạng căn trái.\n• Kẻ viền nét đứt nhẹ cho các ô dữ liệu.",
            "• Dữ liệu hiển thị đẹp mắt, ngày tháng không bị đảo thứ tự tháng/ngày."
        ],
        [
            "5. Căn chỉnh độ rộng tự động\n(Auto-fit Column Width)",
            "Đảm bảo các cột hiển thị vừa khít nội dung, không bị khuất chữ.",
            "• Toàn bộ trang tính sau khi điền đầy đủ dữ liệu.",
            "• Sử dụng phương thức ClosedXML: worksheet.Columns().AdjustToContents().\n• Thiết lập độ rộng tối thiểu cho các cột ghi chú dài.",
            "• Cột Excel tự co giãn vừa vặn, không bị lỗi hiển thị '###' hay mất chữ."
        ],
        [
            "6. Lưu file & Mở tự động\n(Save & Auto-open Prompt)",
            "Lưu file Excel vào máy và hỗ trợ mở ngay cho người dùng xem.",
            "• Mở SaveFileDialog với tên mặc định DanhSachKhachHang_yyyyMMdd.xlsx.",
            "• Lưu file ra đường dẫn người dùng chỉ định.\n• Hiển thị MessageBox: 'Xuất file Excel thành công! Bạn có muốn mở file ngay không?'.\n• Nếu chọn Yes: Gọi Process.Start() mở ngay bằng Microsoft Excel.",
            "• File Excel được lưu trữ an toàn và mở ngay cho người dùng đối soát."
        ]
    ]
    render_function_table(f3_subs)

    doc.add_page_break()

    # =========================================================================
    # PHẦN 2: THÀNH VIÊN 2 - QUẢN LÝ TRÒ CHƠI & DỊCH VỤ
    # =========================================================================
    add_h1("PHẦN 2: CHI TIẾT 6 CHỨC NĂNG NHỎ TRONG CÁC FORM CỦA THÀNH VIÊN 2")

    # Form 4: FormTroChoi
    add_h2("2.1. FORM 4: FormTroChoi (Quản lý Danh mục Trò chơi)")
    add_p("Form phụ trách quản lý danh mục trò chơi trong công viên, theo dõi độ tuổi an toàn, sức chứa mỗi lượt và bảng giá vé lẻ.")

    f4_subs = [
        [
            "1. Tự sinh mã & Thêm trò chơi\n(Auto-ID & Create Game)",
            "Khai báo trò chơi mới vào hệ thống vận hành công viên.",
            "• Tên trò chơi, Loại (Cảm giác mạnh, Trẻ em, Gia đình, Dưới nước, Phiêu lưu), Giá vé, Độ tuổi tối thiểu, Sức chứa, Trạng thái, Mô tả.",
            "• Tự sinh mã TCxxx thông minh (ví dụ: TC016).\n• Validation: Tên không rỗng; Giá vé >= 0; Độ tuổi >= 0; Sức chứa > 0.\n• Gọi GameService.Add(troChoi).",
            "• Lưu bản ghi vào bảng TroChoi.\n• Hiển thị thông báo thành công và cập nhật DataGridView."
        ],
        [
            "2. Xem & Đồng bộ thông tin\n(Select & Bind Game)",
            "Đưa dữ liệu trò chơi lên các điều khiển khi nhấp chọn dòng.",
            "• Sự kiện CellClick trên DataGridView trò chơi.",
            "• Khóa ô txtMaTroChoi (ReadOnly = true).\n• Đổ Tên, Loại, Giá vé, Độ tuổi, Sức chứa, Trạng thái và Mô tả lên form.",
            "• Nhân viên quan sát toàn bộ quy chuẩn kỹ thuật của trò chơi."
        ],
        [
            "3. Cập nhật trò chơi & Quy chuẩn\n(Update Game & Rules)",
            "Thay đổi giá vé lẻ, sức chứa hoặc quy định độ tuổi của trò chơi.",
            "• Thông tin chỉnh sửa trên form nhập liệu.",
            "• Kiểm tra tính hợp lệ dữ liệu.\n• Cập nhật thông tin vào bảng TroChoi qua GameService.Update().",
            "• CSDL được cập nhật, đồng bộ ngay với quầy bán vé POS."
        ],
        [
            "4. Xóa trò chơi có kiểm tra FK\n(Safe Delete Game)",
            "Xóa trò chơi khỏi hệ thống có kiểm tra an toàn chứng từ.",
            "• Nhấn nút [Xóa], MaTroChoi đang chọn.",
            "• Hiển thị cảnh báo xác nhận Yes/No.\n• Kiểm tra ràng buộc khóa ngoại: Truy vấn ChiTietHoaDon xem có dòng nào bán vé trò chơi này không (LoaiHang == 'TroChoi' && MaHang == maTC).\n• Nếu có -> Chặn lại và cảnh báo 'Không thể xóa trò chơi đã phát sinh vé bán!'.",
            "• Xóa an toàn trò chơi nếu chưa từng phát sinh vé bán."
        ],
        [
            "5. Lọc theo Loại & Trạng thái\n(Filter by Category & Status)",
            "Lọc danh mục trò chơi theo phân loại và tình trạng vận hành.",
            "• ComboBox Loại trò chơi và ComboBox Trạng thái (Đang hoạt động / Bảo trì).",
            "• Truy vấn LINQ kết hợp điều kiện lọc kép: LoaiTroChoi == selectedLoai && TrangThai == selectedTrangThai.",
            "• DataGridView hiển thị danh sách trò chơi thỏa mãn bộ lọc."
        ],
        [
            "6. Tìm kiếm trò chơi nhanh\n(Quick Search Game)",
            "Tìm kiếm trò chơi theo tên hoặc mã.",
            "• Từ khóa nhập vào ô txtSearch.",
            "• So khớp Contains không phân biệt hoa thường trên TenTroChoi và MaTroChoi.",
            "• Cập nhật tức thời danh sách trò chơi tìm thấy lên lưới dữ liệu."
        ]
    ]
    render_function_table(f4_subs)

    # Form 5: FormDichVu
    add_h2("2.2. FORM 5: FormDichVu (Quản lý Dịch vụ ăn uống & Tiện ích)")
    add_p("Form quản lý danh mục ẩm thực, giải khát, thuê đồ tiện ích, kiểm soát tồn kho và tự động cảnh báo hết hàng.")

    f5_subs = [
        [
            "1. Tự sinh mã & Thêm dịch vụ\n(Create Park Service)",
            "Thêm mặt hàng ăn uống hoặc dịch vụ tiện ích mới.",
            "• Tên dịch vụ, Phân loại (Ẩm thực, Nước uống, Thuê đồ, Khu nghỉ ngơi, Khác), Đơn giá, Số lượng tồn ban đầu, Mô tả.",
            "• Tự sinh mã DVxxx kế tiếp.\n• Validation: Tên không rỗng; Đơn giá >= 0; Tồn kho >= 0.\n• Tự đặt Trạng thái = 'Đang phục vụ' (nếu tồn > 0) hoặc 'Hết hàng' (nếu tồn = 0).\n• Lưu vào CSDL qua ParkServiceService.Add().",
            "• Bản ghi mới lưu vào bảng DichVu, hiển thị lên DataGridView."
        ],
        [
            "2. Xem & Đồng bộ thông tin\n(Select & Bind Service)",
            "Hiển thị thông tin dịch vụ lên form khi click chọn dòng.",
            "• Sự kiện chọn dòng trên DataGridView.",
            "• Đổ dữ liệu lên TextBox, NumericUpDown, ComboBox; khóa txtMaDichVu.",
            "• Người dùng xem được đơn giá và số lượng tồn hiện tại của mặt hàng."
        ],
        [
            "3. Cập nhật thông tin dịch vụ\n(Update Service Details)",
            "Chỉnh sửa tên mặt hàng, phân loại, đơn giá và mô tả dịch vụ.",
            "• Dữ liệu sửa đổi trên các ô điều khiển.",
            "• Kiểm tra tính hợp lệ dữ liệu, gọi ParkServiceService.Update().",
            "• Dữ liệu dịch vụ được cập nhật chính xác trên toàn hệ thống."
        ],
        [
            "4. Cập nhật nhanh số lượng tồn\n(Quick Stock Update)",
            "Hỗ trợ thủ kho/thu ngân nhập bổ sung số lượng tồn kho nhanh chóng.",
            "• Nhập số lượng mới vào NumericUpDown, bấm [Cập nhật tồn nhanh].",
            "• Kiểm tra số lượng tồn >= 0.\n• Tự động logic: Nếu số lượng = 0 thì tự chuyển TrangThai = 'Hết hàng'; nếu > 0 chuyển 'Đang phục vụ'.\n• Cập nhật trực tiếp trường SoLuongTon vào CSDL.",
            "• Số lượng tồn kho được cập nhật ngay lập tức mà không cần gõ lại các thông tin khác."
        ],
        [
            "5. Xóa dịch vụ an toàn\n(Safe Delete Service)",
            "Xóa mặt hàng dịch vụ nhưng bảo vệ lịch sử hóa đơn kế toán.",
            "• Nhấn nút [Xóa], MaDichVu đang chọn.",
            "• Hộp thoại xác nhận Yes/No.\n• Kiểm tra bảng ChiTietHoaDon (LoaiHang == 'DichVu' && MaHang == maDV). Nếu dịch vụ này đã từng được bán trong bất kỳ hóa đơn nào -> Từ chối xóa để bảo toàn doanh thu.",
            "• Xóa thành công nếu dịch vụ chưa từng bán cho khách."
        ],
        [
            "6. Lọc & Tìm kiếm dịch vụ\n(Filter & Search Service)",
            "Tìm kiếm và lọc dịch vụ theo loại và trạng thái hàng hóa.",
            "• Từ khóa tìm kiếm, ComboBox lọc phân loại (Ẩm thực, Nước uống, Thuê đồ).",
            "• LINQ truy vấn kết hợp lọc theo loại và so khớp tên dịch vụ.",
            "• DataGridView hiển thị danh mục mặt hàng thỏa mãn điều kiện."
        ]
    ]
    render_function_table(f5_subs)

    # Form 6: FormImportExportTroChoiDichVu
    add_h2("2.3. FORM 6: FormImportExportTroChoiDichVu (Nhập/Xuất Excel Trò chơi & Dịch vụ)")
    add_p("Form tích hợp trung tâm TabControl hiện đại hỗ trợ nạp và xuất file Excel cho cả Trò chơi và Dịch vụ.")

    f6_subs = [
        [
            "1. TabControl tích hợp 2 phân hệ\n(Dual-Tab Interface)",
            "Chuyển đổi linh hoạt giữa quản trị Trò chơi và Dịch vụ trong một form.",
            "• Nhấp chuột chọn Tab '🎮 Trò chơi' hoặc Tab '🍔 Dịch vụ'.",
            "• Hiển thị bộ công cụ riêng biệt: Đường dẫn file, Bảng Preview, ListBox lỗi và các nút xử lý độc lập cho từng tab.",
            "• Không gian làm việc trực quan, gọn gàng và chuyên nghiệp."
        ],
        [
            "2. Tải template mẫu theo tab\n(Download Tab Templates)",
            "Xuất sẵn file mẫu Excel chuẩn cột cho Trò chơi hoặc Dịch vụ.",
            "• Bấm nút [Tải file mẫu Trò chơi] hoặc [Tải file mẫu Dịch vụ].",
            "• ClosedXML sinh file TroChoi_Template.xlsx hoặc DichVu_Template.xlsx với các cột quy định và dữ liệu minh họa.",
            "• Người dùng nhận được file mẫu chuẩn cấu trúc để điền dữ liệu."
        ],
        [
            "3. Đọc dữ liệu Excel ClosedXML\n(Excel Parsing Engine)",
            "Đọc toàn bộ các dòng dữ liệu từ file Excel được người dùng chọn.",
            "• Chọn file .xlsx qua nút [Chọn file].",
            "• ClosedXML mở file, quét các ô từ dòng 2 đến LastRowUsed(), trích xuất dữ liệu giá tiền, sức chứa, số lượng tồn kho.",
            "• Tạo danh sách đối tượng thô trong bộ nhớ."
        ],
        [
            "4. Kiểm tra logic nghiệp vụ\n(Business Logic Validation)",
            "Xác thực tính hợp lệ cho từng mặt hàng trước khi lưu.",
            "• Tập dữ liệu thô của từng dòng.",
            "• Với Trò chơi: Giá vé >= 0, Sức chứa > 0, Độ tuổi >= 0; Với Dịch vụ: Đơn giá >= 0, Tồn kho >= 0; Kiểm tra không trùng mã hoặc tên đã có trong CSDL.\n• Thu thập danh sách lỗi chi tiết theo số dòng.",
            "• Phân chia thành ValidItems và ErrorList."
        ],
        [
            "5. Xem trước & Cảnh báo lỗi\n(Preview & Error List)",
            "Hiển thị bảng xem trước dữ liệu hợp lệ và danh sách cảnh báo.",
            "• Kết quả từ bước kiểm tra.",
            "• Đổ ValidItems lên DataGridView Preview.\n• Đổ các dòng lỗi vào ListBox danh sách lỗi kèm lý do cụ thể.\n• Cập nhật nhãn tổng kết số dòng hợp lệ và số dòng lỗi.",
            "• Người dùng kiểm soát chính xác dữ liệu trước khi lưu."
        ],
        [
            "6. Xác nhận nạp & Xuất Excel\n(Confirm Insert & Export)",
            "Ghi dữ liệu hợp lệ vào CSDL hoặc xuất danh sách hiện tại ra Excel.",
            "• Bấm nút [Xác nhận nạp vào CSDL] hoặc nút [Xuất Excel].",
            "• Nạp CSDL: Sinh mã tự động và gọi GameService.AddBatch() / ParkServiceService.AddBatch() trong Transaction.\n• Xuất Excel: ClosedXML tạo file Excel danh mục chuẩn đẹp.",
            "• Dữ liệu được nạp vào SQL Server hoặc xuất ra file báo cáo lưu trữ."
        ]
    ]
    render_function_table(f6_subs)

    doc.add_page_break()

    # =========================================================================
    # PHẦN 3: THÀNH VIÊN 3 - QUẢN LÝ NHÂN VIÊN & BẢNG GIÁ VÉ
    # =========================================================================
    add_h1("PHẦN 3: CHI TIẾT 6 CHỨC NĂNG NHỎ TRONG CÁC FORM CỦA THÀNH VIÊN 3")

    # Form 7: FormNhanVien
    add_h2("3.1. FORM 7: FormNhanVien (Quản lý Hồ sơ Nhân sự)")
    add_p("Form phụ trách quản lý hồ sơ nhân viên trong công viên, chức vụ, trạng thái làm việc và thông tin liên lạc.")

    f7_subs = [
        [
            "1. Tự sinh mã & Thêm nhân viên\n(Auto-ID & Create Employee)",
            "Tiếp nhận hồ sơ nhân viên mới vào làm việc tại công viên.",
            "• Họ tên, Ngày sinh, Giới tính, Số điện thoại, Email, Địa chỉ, Chức vụ (Quản lý, Thu ngân, Soát vé, Kỹ thuật, Bảo vệ, Chăm sóc KH), Ngày vào làm, Trạng thái.",
            "• Tự sinh mã NVxxx thông minh (ví dụ: NV011).\n• Validation: Bắt buộc Họ tên và SĐT; SĐT đủ 10 số; Email đúng định dạng; Kiểm tra trùng SĐT trong CSDL.\n• Gọi EmployeeService.Add(nhanVien).",
            "• Thêm bản ghi mới vào bảng NhanVien.\n• Thông báo thành công và làm mới DataGridView."
        ],
        [
            "2. Xem & Đồng bộ hồ sơ nhân sự\n(Select & Bind Employee)",
            "Đưa dữ liệu nhân viên lên các ô điều khiển khi chọn dòng.",
            "• Click chọn dòng nhân viên trên DataGridView.",
            "• Khóa txtMaNV (ReadOnly = true).\n• Đổ toàn bộ thông tin nhân sự lên TextBox, DateTimePicker, ComboBox.",
            "• Xem chi tiết lý lịch và chức vụ hiện tại của nhân viên."
        ],
        [
            "3. Cập nhật hồ sơ & Chức vụ\n(Update Employee Profile)",
            "Điều chuyển chức vụ, thay đổi thông tin liên lạc hoặc trạng thái làm việc.",
            "• Dữ liệu đã thay đổi trên form.",
            "• Kiểm tra tính hợp lệ dữ liệu và kiểm tra không trùng SĐT với nhân viên khác.\n• Cập nhật CSDL qua EmployeeService.Update().",
            "• Hồ sơ nhân viên được cập nhật chính xác trên toàn hệ thống."
        ],
        [
            "4. Xóa nhân viên an toàn\n(Safe Delete Employee)",
            "Xóa hồ sơ nhân sự nhưng bảo toàn chứng từ thu ngân.",
            "• Nhấn nút [Xóa], MaNV của nhân viên đang chọn.",
            "• Hộp thoại xác nhận Yes/No.\n• Kiểm tra ràng buộc khóa ngoại: Truy vấn bảng HoaDon. Nếu nhân viên này đã từng lập hóa đơn bán vé (existing.HoaDons.Any()) -> Chặn lại và cảnh báo 'Không thể xóa nhân viên đã từng lập hóa đơn!'.\n• Nếu chưa có hóa đơn: Cho phép xóa.",
            "• Bảo vệ toàn vẹn lịch sử thu ngân, không vi phạm khóa ngoại."
        ],
        [
            "5. Lọc theo Chức vụ & Trạng thái\n(Filter by Role & Status)",
            "Lọc danh sách nhân viên theo phòng ban/vị trí công tác.",
            "• ComboBox Chức vụ và ComboBox Trạng thái (Đang làm việc / Đã nghỉ việc).",
            "• LINQ truy vấn kết hợp điều kiện lọc ChucVu và TrangThai.",
            "• DataGridView chỉ hiển thị các nhân viên phù hợp tiêu chí lọc."
        ],
        [
            "6. Tìm kiếm nhân viên nhanh\n(Quick Search Employee)",
            "Tìm kiếm nhân sự theo từ khóa mã, họ tên hoặc số điện thoại.",
            "• Nhập từ khóa vào ô txtSearch.",
            "• So khớp Contains trên MaNV, HoTen, SoDienThoai, Email.",
            "• Hiển thị danh sách nhân sự tìm thấy lên lưới dữ liệu."
        ]
    ]
    render_function_table(f7_subs)

    # Form 8: FormVe
    add_h2("3.2. FORM 8: FormVe (Quản lý Bảng giá các Loại Vé)")
    add_p("Form quản lý bảng giá các loại vé vào cổng, vé trọn gói, vé VIP và thời hạn áp dụng vé.")

    f8_subs = [
        [
            "1. Tự sinh mã & Thêm loại vé\n(Auto-ID & Create Ticket)",
            "Thiết lập loại vé mới phục vụ khách tham quan công viên.",
            "• Tên loại vé, Phân loại (Vé ngày, Vé trẻ em, Vé người lớn, Vé gia đình, Vé VIP), Giá vé, Thời hạn (Trong ngày, 1 tháng...), Mô tả quyền lợi, Trạng thái.",
            "• Tự sinh mã VExxx kế tiếp (ví dụ: VE011).\n• Validation: Tên vé không rỗng; Giá vé >= 0; Thời hạn không rỗng.\n• Gọi TicketService.Add(ve).",
            "• Bản ghi mới được lưu vào bảng Ve.\n• Cập nhật DataGridView bảng giá vé."
        ],
        [
            "2. Xem & Đồng bộ thông tin vé\n(Select & Bind Ticket)",
            "Hiển thị thông tin loại vé lên các ô nhập liệu khi chọn dòng.",
            "• Sự kiện chọn dòng trên DataGridView.",
            "• Khóa txtMaVe (ReadOnly = true).\n• Đổ Tên vé, Loại vé, Giá vé, Thời hạn, Trạng thái và Mô tả lên form.",
            "• Quan sát quyền lợi và giá vé của loại vé đang chọn."
        ],
        [
            "3. Điều chỉnh bảng giá vé\n(Update Ticket Price)",
            "Cập nhật giá vé mới hoặc sửa đổi mô tả quyền lợi của vé.",
            "• Thông tin chỉnh sửa trên form.",
            "• Kiểm tra tính hợp lệ dữ liệu, Giá vé >= 0.\n• Cập nhật thông tin vào bảng Ve qua TicketService.Update().",
            "• Giá vé mới có hiệu lực ngay lập tức tại quầy bán vé POS."
        ],
        [
            "4. Đổi trạng thái Áp dụng\n(Toggle Active Status)",
            "Bật hoặc tắt trạng thái áp dụng loại vé khi hết mùa ưu đãi.",
            "• Chuyển đổi ComboBox Trạng thái: 'Đang áp dụng' <-> 'Ngừng áp dụng'.",
            "• Cập nhật trường TrangThai vào CSDL.",
            "• Vé bị 'Ngừng áp dụng' sẽ tự động ẩn khỏi màn hình bán vé POS, tránh thu ngân bán nhầm."
        ],
        [
            "5. Xóa loại vé có kiểm tra FK\n(Safe Delete Ticket)",
            "Xóa loại vé có kiểm tra ràng buộc lịch sử giao dịch.",
            "• Nhấn nút [Xóa], MaVe đang chọn.",
            "• Cảnh báo xác nhận Yes/No.\n• Kiểm tra bảng ChiTietHoaDon (LoaiHang == 'Ve' && MaHang == maVe). Nếu loại vé này đã từng được bán trong hóa đơn -> Chặn lại và hướng dẫn chuyển sang 'Ngừng áp dụng'.",
            "• Bảo vệ toàn vẹn dữ liệu bán hàng lịch sử của công viên."
        ],
        [
            "6. Lọc & Tìm kiếm bảng giá vé\n(Filter & Search Tickets)",
            "Tìm kiếm và lọc vé theo phân loại đối tượng sử dụng.",
            "• Từ khóa tìm kiếm, ComboBox lọc phân loại (Vé trẻ em, người lớn, VIP).",
            "• LINQ truy vấn lọc theo loại vé và so khớp tên vé.",
            "• DataGridView hiển thị danh mục vé theo đúng yêu cầu."
        ]
    ]
    render_function_table(f8_subs)

    # Form 9: FormImportExportNhanVienVe
    add_h2("3.3. FORM 9: FormImportExportNhanVienVe (Nhập/Xuất Excel Nhân viên & Vé)")
    add_p("Form tích hợp TabControl quản lý nhập xuất file Excel cho cả 2 danh mục Nhân sự và Bảng giá vé.")

    f9_subs = [
        [
            "1. TabControl tích hợp 2 phân hệ\n(Dual-Tab Employee & Ticket)",
            "Chuyển đổi linh hoạt giữa quản trị Nhân viên và Loại vé trong một form.",
            "• Nhấp chọn Tab '👥 Quản lý Nhân viên' hoặc Tab '🎫 Quản lý Loại vé'.",
            "• Hiển thị bộ công cụ riêng biệt cho từng tab với đường dẫn file, Preview, ListBox lỗi và nút bấm độc lập.",
            "• Giao diện khoa học, đồng bộ và tiện ích cho người quản lý."
        ],
        [
            "2. Tải template mẫu theo tab\n(Download Templates)",
            "Xuất file Excel mẫu chuẩn cột cho Nhân viên hoặc Bảng giá vé.",
            "• Bấm [Tải file mẫu Nhân viên] hoặc [Tải file mẫu Vé].",
            "• ClosedXML sinh file NhanVien_Template.xlsx hoặc Ve_Template.xlsx kèm hàng mẫu minh họa.",
            "• Người dùng có ngay file chuẩn để điền dữ liệu."
        ],
        [
            "3. Đọc dữ liệu Excel ClosedXML\n(Excel Parsing Engine)",
            "Đọc toàn bộ các hàng dữ liệu từ file Excel được chọn.",
            "• Chọn file .xlsx từ máy tính.",
            "• ClosedXML mở file, duyệt từ dòng 2 đến dòng cuối cùng, bóc tách dữ liệu chuỗi, số tiền và ngày tháng.",
            "• Nạp danh sách thô vào bộ nhớ phục vụ kiểm duyệt."
        ],
        [
            "4. Kiểm tra toàn vẹn nghiệp vụ\n(Row Validation Engine)",
            "Xác thực tính hợp lệ cho từng nhân viên hoặc loại vé.",
            "• Tập dữ liệu thô từng dòng.",
            "• Với Nhân viên: SĐT đủ 10 số, email chuẩn, ngày vào làm hợp lệ; Với Vé: Giá vé >= 0, tên vé bắt buộc; Kiểm tra không trùng mã hoặc SĐT đã có trong CSDL.\n• Ghi lại lỗi chi tiết kèm vị trí dòng.",
            "• Tách thành ValidItems và ErrorList."
        ],
        [
            "5. Xem trước & Cảnh báo lỗi\n(Preview & Error List)",
            "Hiển thị bảng xem trước dữ liệu hợp lệ và danh sách lỗi chi tiết.",
            "• Kết quả từ bước kiểm tra.",
            "• Đổ ValidItems lên DataGridView Preview.\n• Đổ các dòng lỗi vào ListBox lstErrors kèm nguyên nhân cụ thể.\n• Cập nhật nhãn tổng kết số lượng dòng hợp lệ và lỗi.",
            "• Người dùng an tâm kiểm tra trước khi ghi dữ liệu vào CSDL."
        ],
        [
            "6. Xác nhận nạp & Xuất Excel\n(Confirm Insert & Export)",
            "Ghi dữ liệu hợp lệ vào CSDL hoặc xuất danh sách hiện hành ra Excel.",
            "• Bấm nút [Xác nhận nạp vào CSDL] hoặc nút [Xuất Excel].",
            "• Nạp CSDL: Sinh mã kế tiếp tự động và gọi EmployeeService.AddBatch() / TicketService.AddBatch() trong Transaction.\n• Xuất Excel: ClosedXML tạo file Excel danh mục chuẩn đẹp.",
            "• Dữ liệu được nạp vào SQL Server hoặc xuất ra file báo cáo lưu trữ."
        ]
    ]
    render_function_table(f9_subs)

    doc.add_page_break()

    # =========================================================================
    # PHẦN 4: THÀNH VIÊN 4 - BÁN VÉ POS, HÓA ĐƠN & BÁO CÁO DOANH THU
    # =========================================================================
    add_h1("PHẦN 4: CHI TIẾT 6 CHỨC NĂNG NHỎ TRONG CÁC FORM CỦA THÀNH VIÊN 4")

    # Form 10: FormBanVe
    add_h2("4.1. FORM 10: FormBanVe (Quầy Tạo đơn Bán vé POS)")
    add_p("Màn hình quầy bán vé (Point of Sale) tốc độ cao phục vụ khách mua vé vào cổng, trò chơi và dịch vụ ăn uống kèm quản lý tồn kho tức thời.")

    f10_subs = [
        [
            "1. Danh mục Card sản phẩm\n(Interactive Product Cards)",
            "Hiển thị danh mục mặt hàng trực quan phân chia theo nhóm thẻ Card.",
            "• Chọn Tab nhóm hàng: 'Tất cả', 'Vé vào cổng', 'Trò chơi', 'Dịch vụ'.",
            "• Đọc danh mục từ CSDL, sinh động các thẻ Panel sản phẩm có Tên, Đơn giá, Số lượng tồn kho (nếu là dịch vụ) và nút [+ Thêm].",
            "• Giao diện POS hiện đại, dễ thao tác nhanh trên màn hình cảm ứng hoặc chuột."
        ],
        [
            "2. Thêm vào giỏ & Kiểm soát tồn\n(Add to Cart & Stock Check)",
            "Thêm mặt hàng vào giỏ hàng và kiểm soát tồn kho tức thời.",
            "• Nhấn nút [+ Thêm] trên thẻ sản phẩm.",
            "• Kiểm tra nếu là Dịch vụ: Kiểm tra số lượng yêu cầu so với SoLuongTon hiện tại. Nếu vượt tồn -> Cảnh báo 'Hàng trong kho không đủ!'. Nếu hợp lệ -> Thêm vào List<CartItemDto> hoặc tăng số lượng nếu đã có trong giỏ.",
            "• Mặt hàng xuất hiện trong danh sách giỏ hàng bên phải."
        ],
        [
            "3. Điều chỉnh số lượng & Xóa giỏ\n(Cart Item Modification)",
            "Tăng/giảm số lượng hoặc xóa mặt hàng trong giỏ hàng.",
            "• Bấm nút [+] tăng, nút [-] giảm số lượng, nút [🗑️ Xóa dòng], hoặc nút [Xóa tất cả].",
            "• Cập nhật số lượng mặt hàng, kiểm tra giới hạn tồn kho, tự tính lại thành tiền từng dòng: ThanhTien = SoLuong * DonGia.",
            "• Giỏ hàng cập nhật số lượng tức thời."
        ],
        [
            "4. Tính tổng tiền & Chiết khấu\n(Real-time Totals Calculation)",
            "Tự động tính Tạm tính, Chiết khấu và Tổng tiền thanh toán.",
            "• Thay đổi số lượng mặt hàng hoặc nhập phần trăm giảm giá (0 - 100%).",
            "• Công thức: TamTinh = Sum(ThanhTien); TienGiam = TamTinh * (GiamGia / 100); TongTien = TamTinh - TienGiam.\n• Định dạng hiển thị tiền tệ VNĐ (#,##0 VNĐ).",
            "• Các nhãn lblSubtotal và lblTotal cập nhật số tiền chính xác theo thời gian thực."
        ],
        [
            "5. Chọn khách hàng & Thanh toán\n(Customer & Payment Method)",
            "Gán khách hàng thành viên và phương thức thanh toán cho hóa đơn.",
            "• Chọn khách hàng từ ComboBox (hoặc để trống nếu khách vãng lai), chọn Phương thức (Tiền mặt, Chuyển khoản QR, Quẹt thẻ), nhập ghi chú.",
            "• Lưu thông tin MaKH, MaNV (nhân viên đăng nhập), phương thức thanh toán vào đối tượng đơn hàng.",
            "• Hoàn tất đầy đủ thông tin pháp lý cho hóa đơn bán hàng."
        ],
        [
            "6. Thanh toán Transaction & In HĐ\n(Checkout Transaction & Print)",
            "Thực thi thanh toán an toàn, trừ tồn kho và hiển thị hóa đơn in nhiệt.",
            "• Nhấn nút lớn [THANH TOÁN & IN HÓA ĐƠN].",
            "• Mở Database Transaction:\n  + Sinh mã hóa đơn mới HDxxx.\n  + Lưu bản ghi HoaDon.\n  + Lưu các bản ghi ChiTietHoaDon.\n  + Tự động trừ tồn kho mặt hàng Dịch vụ (dv.SoLuongTon -= item.SoLuong). Nếu tồn về 0 chuyển TrangThai = 'Hết hàng'.\n  + Commit Transaction. Nếu có lỗi -> Rollback ngay lập tức.\n• Hiển thị FormChiTietHoaDon xem trước hóa đơn để in. Làm sạch giỏ hàng.",
            "• Giao dịch thành công, CSDL toàn vẹn, in phiếu thu cho khách."
        ]
    ]
    render_function_table(f10_subs)

    # Form 11: FormHoaDon
    add_h2("4.2. FORM 11: FormHoaDon (Quản lý Danh sách Hóa đơn & Lịch sử Giao dịch)")
    add_p("Form quản lý toàn bộ các hóa đơn đã giao dịch, hỗ trợ tra cứu chi tiết và thực thi quy trình hủy hóa đơn an toàn có hoàn trả kho hàng.")

    f11_subs = [
        [
            "1. Lọc hóa đơn theo khoảng ngày\n(Date Range Filtering)",
            "Tra cứu các giao dịch phát sinh trong khoảng thời gian xác định.",
            "• Chọn 'Từ ngày' (dtpFrom) và 'Đến ngày' (dtpTo).",
            "• LINQ truy vấn hóa đơn có NgayLap.Date nằm trong khoảng từ dtpFrom đến dtpTo, sắp xếp giảm dần theo ngày lập mới nhất.",
            "• DataGridView hiển thị danh sách hóa đơn trong kỳ."
        ],
        [
            "2. Lọc theo Khách hàng & Từ khóa\n(Customer Filter & Search)",
            "Lọc hóa đơn theo khách hàng cụ thể hoặc tìm theo mã hóa đơn, thu ngân.",
            "• Chọn khách hàng từ ComboBox hoặc nhập từ khóa tìm kiếm vào txtSearch.",
            "• Truy vấn kết hợp điều kiện lọc: MaKH, MaHoaDon, hoặc tên nhân viên lập đơn.",
            "• Lưới dữ liệu hiển thị các hóa đơn khớp tiêu chí tìm kiếm."
        ],
        [
            "3. Xem chi tiết phiếu thanh toán\n(View Invoice Detail Modal)",
            "Mở cửa sổ xem lại chi tiết đầy đủ của hóa đơn đã chọn.",
            "• Chọn một hóa đơn trên bảng và nhấn nút [Xem chi tiết].",
            "• Khởi tạo FormChiTietHoaDon truyền vào MaHoaDon đang chọn; nạp toàn bộ danh mục vé, dịch vụ, số lượng, đơn giá và thành tiền.",
            "• Hiển thị phiếu thanh toán chuyên nghiệp, sẵn sàng in lại cho khách."
        ],
        [
            "4. Hủy hóa đơn an toàn & Hoàn kho\n(Safe Cancel Invoice & Return)",
            "Hủy giao dịch khi khách trả vé và tự động hoàn trả tồn kho.",
            "• Chọn hóa đơn và nhấn nút [Hủy hóa đơn].",
            "• Kiểm tra: Nếu hóa đơn đã 'Đã hủy' -> Từ chối; Nhập lý do hủy đơn. Trong Database Transaction:\n  + Đổi TrangThai = 'Đã hủy'.\n  + Ghi lý do hủy vào GhiChu.\n  + Duyệt các dòng ChiTietHoaDon là Dịch vụ: Tự động cộng hoàn trả lại số lượng tồn kho (dv.SoLuongTon += item.SoLuong) và khôi phục trạng thái 'Đang phục vụ'.\n  + Commit Transaction.",
            "• Hóa đơn được hủy an toàn, kho hàng tự động cân bằng chính xác với thực tế."
        ],
        [
            "5. Tổng kết số lượng & Doanh thu\n(Summary Calculation)",
            "Tính toán tổng số hóa đơn và tổng tiền thực thu trên danh sách đang hiển thị.",
            "• Bất kỳ thay đổi nào trong bộ lọc hoặc tìm kiếm.",
            "• Tính tổng số hóa đơn tìm thấy và tổng tiền thực thu: Sum(TongTien) của các hóa đơn có trạng thái 'Đã thanh toán'.",
            "• Cập nhật các nhãn lblTotal và lblTotalAmount trực quan trên thanh tiêu đề."
        ],
        [
            "6. Xuất danh sách hóa đơn ra Excel\n(Export Invoice List)",
            "Xuất toàn bộ danh sách hóa đơn đang lọc ra file Excel bằng ClosedXML.",
            "• Nhấn nút [Xuất Excel].",
            "• ClosedXML tạo file Excel danh sách hóa đơn (Mã HĐ, Khách hàng, Thu ngân, Ngày lập, Giảm giá %, Tổng tiền, Phương thức thanh toán, Trạng thái), định dạng tiền tệ và độ rộng cột tự động.",
            "• File Excel được tạo phục vụ công tác kế toán và đối soát cuối ngày."
        ]
    ]
    render_function_table(f11_subs)

    # Form 12: FormBaoCaoDoanhThu
    add_h2("4.3. FORM 12: FormBaoCaoDoanhThu (Báo cáo Thống kê Doanh thu & Cơ cấu)")
    add_p("Form cung cấp báo cáo tài chính toàn diện theo kỳ, phân tích biến động dòng tiền và cơ cấu nguồn thu theo từng mảng kinh doanh.")

    f12_subs = [
        [
            "1. Lọc kỳ báo cáo linh hoạt\n(Period Selection Engine)",
            "Thiết lập mốc thời gian báo cáo tài chính nhanh chóng.",
            "• Chọn ComboBox: 'Hôm nay', 'Tuần này', 'Tháng này', 'Năm nay', 'Tùy chọn'.",
            "• Tự động tính toán mốc ngày bắt đầu và kết thúc (DateTime.Today, đầu tuần, đầu tháng, đầu năm) và cập nhật hai DateTimePicker dtpFrom, dtpTo.",
            "• Xác định chính xác khung thời gian cần tổng hợp số liệu kế toán."
        ],
        [
            "2. Tổng hợp 3 chỉ số tài chính KPI\n(Summary KPI Cards)",
            "Tính toán 3 chỉ số tài chính cốt lõi của công viên trong kỳ.",
            "• Khoảng thời gian báo cáo được chọn.",
            "• Truy vấn các hóa đơn có trạng thái 'Đã thanh toán':\n  + Tổng doanh thu = Sum(TongTien).\n  + Tổng số vé đã bán = Sum(SoLuong) của các dòng có LoaiHang == 'Ve'.\n  + Tổng số đơn = Count().",
            "• Hiển thị nổi bật trên 3 thẻ Card thống kê màu sắc sang trọng."
        ],
        [
            "3. Bảng thống kê theo thời gian\n(Timeline Revenue Breakdown)",
            "Phân tích diễn biến doanh thu theo từng mốc thời gian chi tiết.",
            "• Kỳ báo cáo đã chọn.",
            "• Nhóm (GroupBy) hóa đơn theo mốc thời gian (từng ngày trong tháng hoặc từng tháng trong năm), tính tổng doanh thu, số vé bán, số đơn của từng mốc.",
            "• DataGridView 1 hiển thị chi tiết biến động doanh thu theo dòng thời gian."
        ],
        [
            "4. Phân tích cơ cấu nguồn thu\n(Revenue by Category)",
            "Phân tích cơ cấu tỷ trọng doanh thu theo 3 nhóm mặt hàng chính.",
            "• Kỳ báo cáo đã chọn.",
            "• Phân rã doanh thu theo 3 nhóm: 'Vé vào cổng (Ve)', 'Dịch vụ tiện ích (DichVu)', 'Vé trò chơi lẻ (TroChoi)'. Tính tổng tiền, số lượng bán ra và tỷ lệ phần trăm đóng góp doanh thu của từng nhóm.",
            "• DataGridView 2 hiển thị cơ cấu nguồn thu, giúp nhà quản lý nắm rõ mặt hàng nào mang lại lợi nhuận cao nhất."
        ],
        [
            "5. Lọc danh mục cơ cấu tỷ trọng\n(Category Filter)",
            "Lọc nhanh bảng cơ cấu doanh thu theo nhóm mặt hàng mong muốn.",
            "• ComboBox phân loại: 'Tất cả', 'Chỉ xem Vé', 'Chỉ xem Dịch vụ', 'Chỉ xem Trò chơi'.",
            "• Lọc nhanh bảng cơ cấu doanh thu theo nhóm được chọn.",
            "• Cung cấp góc nhìn chuyên sâu cho từng mảng kinh doanh của công viên."
        ],
        [
            "6. Xuất báo cáo tài chính ra Excel\n(Export Financial Report)",
            "Xuất báo cáo tài chính tổng hợp đa bảng ra file Excel ClosedXML.",
            "• Bấm nút [Xuất báo cáo Doanh thu ra Excel].",
            "• ClosedXML tạo file báo cáo tài chính đa bảng: Gồm banner tiêu đề công viên, khoảng thời gian báo cáo, các chỉ số KPI, bảng diễn biến doanh thu theo thời gian và bảng cơ cấu tỷ trọng mặt hàng. Có định dạng dòng Total đậm nét và số tiền định dạng tiền tệ VNĐ.",
            "• File Excel báo cáo tài chính hoàn chỉnh để nộp cho Ban Giám đốc hoặc Giảng viên chấm điểm."
        ]
    ]
    render_function_table(f12_subs)

    doc.add_page_break()

    # =========================================================================
    # TỔNG KẾT & ĐỐI CHIẾU TIÊU CHÍ ĐÁNH GIÁ
    # =========================================================================
    add_h1("TỔNG KẾT VÀ ĐỐI CHIẾU TIÊU CHÍ ĐÁNH GIÁ BÀI TẬP LỚN")
    add_p("Hệ thống phần mềm Quản lý Khu vui chơi Wonderland Park (FunZone) với 72 chức năng con phân bổ trong 12 Forms lớn đáp ứng xuất sắc các tiêu chí chấm điểm chuyên môn:")
    add_bullet("Tính toàn vẹn và chuẩn hóa CSDL: ", "Cấu trúc 8 bảng chuẩn 3NF, thiết lập đầy đủ khóa chính, khóa ngoại, chỉ mục tìm kiếm và ràng buộc kiểm tra toàn vẹn tham chiếu. Không cho phép xóa dữ liệu cha khi đã phát sinh hóa đơn bán hàng.")
    add_bullet("Tính an toàn giao dịch (Transactions): ", "Các hành động tạo đơn bán vé (POS) và hủy hóa đơn (Cancel Invoice) được bao bọc 100% trong Database Transaction, bảo đảm cập nhật đồng thời chứng từ và số lượng tồn kho dịch vụ an toàn tuyệt đối.")
    add_bullet("Khả năng tích hợp Excel chuyên nghiệp: ", "Sử dụng ClosedXML xử lý đọc ghi file .xlsx hai chiều, có template mẫu, pipeline kiểm lỗi từng dòng có báo cáo chi tiết và định dạng bảng báo cáo chuẩn doanh nghiệp.")
    add_bullet("Trải nghiệm người dùng (UX/UI): ", "Giao diện Windows Forms hiện đại, Flat Design, màu sắc hài hòa, hỗ trợ nhúng Form mượt mà trên FormMain, tối ưu hóa thao tác cho nhân viên bán vé và nhà quản trị.")

    # Save to both paths
    out_dir_1 = r"c:\Users\KIM THIEN\OneDrive\Desktop\Playground\QuanLyKhuVuiChoi"
    out_file_1 = os.path.join(out_dir_1, "Chi_Tiet_6_Chuc_Nang_Nho_12_Form.docx")
    doc.save(out_file_1)
    print(f"Saved successfully to: {out_file_1}")

    out_dir_2 = r"c:\Users\KIM THIEN\OneDrive\Desktop\Playground"
    out_file_2 = os.path.join(out_dir_2, "Chi_Tiet_6_Chuc_Nang_Nho_12_Form.docx")
    doc.save(out_file_2)
    print(f"Saved successfully to: {out_file_2}")

if __name__ == "__main__":
    create_document()
