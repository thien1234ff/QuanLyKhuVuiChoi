import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Set page margins (Normal: 1 inch = 72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Base colors
    COLOR_PRIMARY = RGBColor(31, 78, 120)     # #1F4E78 Dark Navy
    COLOR_SECONDARY = RGBColor(46, 117, 182)  # #2E75B6 Blue
    COLOR_ACCENT = RGBColor(16, 185, 129)     # #10B981 Emerald
    COLOR_TEXT = RGBColor(30, 41, 59)         # #1E293B Slate Dark
    COLOR_MUTED = RGBColor(100, 116, 139)     # #64748B Slate Muted

    # Setup normal style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(4)

    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D3D3D3"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

    def add_callout(text, title=None, alert_type="note"):
        # Alert background and border colors
        bg_color = "EFF6FF"  # light blue
        border_color = "2563EB" # blue
        icon = "📌"
        if alert_type == "success":
            bg_color = "ECFDF5"
            border_color = "10B981"
            icon = "✅"
        elif alert_type == "warning":
            bg_color = "FFFBEB"
            border_color = "F59E0B"
            icon = "⚠️"

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.7)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        # left border only
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
                <w:top w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if title:
            r_title = p.add_run(f"{icon} {title}\n")
            r_title.bold = True
            r_title.font.name = "Segoe UI"
            r_title.font.size = Pt(10.5)
            r_title.font.color.rgb = RGBColor(30, 58, 138) if alert_type == "note" else RGBColor(6, 95, 70)
        r_text = p.add_run(text)
        r_text.font.name = "Segoe UI"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_TEXT

        # add a small spacer paragraph after table
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Segoe UI'
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10.5)
        r.font.italic = italic
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Segoe UI'
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_TEXT
        return p

    def format_table(table, col_widths, headers, data, align_cols=None):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].width = Inches(col_widths[i])
            set_cell_background(hdr_cells[i], "1F4E78")
            set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
            p = hdr_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h)
            r.bold = True
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_data in enumerate(data):
            row_cells = table.rows[r_idx + 1].cells
            bg_col = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].width = Inches(col_widths[c_idx])
                set_cell_background(row_cells[c_idx], bg_col)
                set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if align_cols and c_idx in align_cols:
                    p.alignment = align_cols[c_idx]
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(val))
                r.font.name = 'Segoe UI'
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_TEXT

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 0. TRANG TIÊU ĐỀ (COVER BANNER)
    # =========================================================================
    p_title_pre = doc.add_paragraph()
    p_title_pre.paragraph_format.space_before = Pt(20)
    p_title_pre.paragraph_format.space_after = Pt(4)
    p_title_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_title_pre.add_run("BÁO CÁO BÀI TẬP LỚN MÔN LẬP TRÌNH CƠ SỞ DỮ LIỆU\n")
    r_uni.font.name = "Segoe UI"
    r_uni.font.size = Pt(12)
    r_uni.bold = True
    r_uni.font.color.rgb = COLOR_SECONDARY

    p_main_title = doc.add_paragraph()
    p_main_title.paragraph_format.space_before = Pt(10)
    p_main_title.paragraph_format.space_after = Pt(8)
    p_main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_main_title.add_run("TÀI LIỆU ĐẶC TẢ CHI TIẾT CÁC CHỨC NĂNG HỆ THỐNG\nQUẢN LÝ KHU VUI CHƠI GIẢI TRÍ (WONDERLAND PARK / FUNZONE)")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(20)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Kiến trúc 3-Layer | C# .NET 8 WinForms | Microsoft SQL Server | Entity Framework Core 8 | ClosedXML")
    r_sub.italic = True
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = COLOR_MUTED

    # Box tóm tắt thông tin dự án
    add_callout(
        "• Tên đề tài: Hệ thống phần mềm Quản lý Khu vui chơi Giải trí Wonderland Park (FunZone).\n"
        "• Công nghệ: .NET 8.0 Windows Forms, Entity Framework Core 8.0, Microsoft SQL Server, ClosedXML Excel.\n"
        "• Quy mô: 12 Forms chuẩn nghiệp vụ, 8 bảng CSDL quan hệ chặt chẽ, hỗ trợ đầy đủ 6 thao tác bắt buộc (Thêm, Sửa, Xóa, Tìm kiếm, Import Excel, Export Excel).\n"
        "• Phân chia trách nhiệm: 4 thành viên đảm nhiệm trọn vẹn từ Giao diện (Presentation), Xử lý nghiệp vụ (BLL) đến Truy xuất dữ liệu (DAL).",
        title="THÔNG TIN TỔNG QUAN HỆ THỐNG",
        alert_type="note"
    )

    doc.add_page_break()

    # =========================================================================
    # 1. TỔNG QUAN DỰ ÁN
    # =========================================================================
    add_h1("1. TỔNG QUAN DỰ ÁN VÀ MỤC TIÊU HỆ THỐNG")
    
    add_h2("1.1. Bối cảnh và Mục tiêu")
    add_p("Trong thời đại chuyển đổi số, việc vận hành các khu vui chơi giải trí phức hợp (quy mô từ công viên nước, trò chơi mạo hiểm, trò chơi gia đình đến chuỗi dịch vụ ăn uống, cho thuê phụ kiện) đòi hỏi sự quản lý chặt chẽ, tức thời và chính xác. Các phương thức bán vé thủ công hoặc ghi chép sổ sách truyền thống thường xuyên bộc lộ hạn chế: thất thoát doanh thu, thời gian khách chờ đợi lâu, khó kiểm soát tồn kho dịch vụ và thiếu dữ liệu thống kê phục vụ chiến lược kinh doanh.")
    add_p("Hệ thống phần mềm 'Quản lý Khu vui chơi Giải trí Wonderland Park (FunZone)' được xây dựng nhằm tin học hóa toàn diện quy trình vận hành của một khu vui chơi hiện đại:")
    add_bullet("Tự động hóa quầy bán vé (POS): ", "Tạo đơn bán vé, dịch vụ và combo trò chơi chỉ trong vài giây; hỗ trợ tính toán chiết khấu, nhiều hình thức thanh toán (Tiền mặt, Chuyển khoản, Quẹt thẻ) và tự động trừ tồn kho mặt hàng.")
    add_bullet("Quản lý khách hàng thân thiết: ", "Lưu trữ lịch sử giao dịch, hỗ trợ tìm kiếm linh hoạt đa phương thức và chăm sóc khách hàng chuyên nghiệp.")
    add_bullet("Kiểm soát hạ tầng trò chơi & dịch vụ: ", "Theo dõi trạng thái vận hành, bảo trì, sức chứa an toàn, độ tuổi giới hạn của trò chơi; giám sát số lượng tồn kho dịch vụ ăn uống và đồ lưu niệm.")
    add_bullet("Quản trị nhân sự & ca làm việc: ", "Phân quyền chi tiết giữa Quản trị viên (Admin) và Nhân viên thu ngân/soát vé, lưu trữ hồ sơ công tác.")
    add_bullet("Báo cáo thông minh & Tích hợp Excel: ", "Xuất - nhập dữ liệu nhanh chóng qua file Excel định dạng chuẩn quốc tế bằng ClosedXML, cung cấp dashboard biểu đồ doanh thu theo thời gian thực.")

    add_h2("1.2. Đối tượng sử dụng và Phân quyền")
    table_role = doc.add_table(rows=3, cols=3)
    format_table(
        table_role,
        [1.3, 1.8, 3.6],
        ["Vai trò", "Đối tượng", "Quyền hạn & Phạm vi truy cập trên phần mềm"],
        [
            ["Quản trị viên (Admin)", "Ban Giám đốc, Quản lý điều hành", "Toàn quyền trên toàn hệ thống: Xem Dashboard tổng quan; Quản lý khách hàng; Thêm/sửa/xóa Trò chơi, Dịch vụ; Quản lý hồ sơ Nhân viên; Thiết lập bảng giá Vé; Xem báo cáo doanh thu chi tiết; Nhập/Xuất Excel toàn bộ các danh mục."],
            ["Thu ngân / Nhân viên (NhanVien)", "Nhân viên quầy bán vé (POS), Quầy dịch vụ", "Truy cập màn hình Bán vé (POS), tra cứu thông tin khách hàng, xem danh sách hóa đơn bán trong ca, tìm kiếm bảng giá vé và danh mục trò chơi/dịch vụ còn phục vụ. Bị giới hạn chỉnh sửa hồ sơ nhân viên và báo cáo tài chính cấp cao."]
        ]
    )

    # =========================================================================
    # 2. KIẾN TRÚC HỆ THỐNG VÀ CÔNG NGHỆ ÁP DỤNG
    # =========================================================================
    add_h1("2. KIẾN TRÚC HỆ THỐNG VÀ CÔNG NGHỆ ÁP DỤNG")
    
    add_h2("2.1. Ngôn ngữ & Công nghệ sử dụng")
    add_bullet("Ngôn ngữ lập trình: ", "C# phiên bản 12 mới nhất, tận dụng các cú pháp hiện đại: Record, Pattern Matching, Tuple returns, Nullable Reference Types.")
    add_bullet("Nền tảng ứng dụng: ", ".NET 8.0 (net8.0-windows) với Windows Forms hiện đại hóa. Giao diện được thiết kế theo phong cách Flat Design, tông màu hài hòa (Navy #1F4E78, Royal Blue #3B66F5, Emerald #10B981), bo góc mềm mại, tối ưu trải nghiệm người dùng (UX/UI).")
    add_bullet("Hệ quản trị CSDL: ", "Microsoft SQL Server 2019/2022 / SQL Server LocalDB. Hỗ trợ đầy đủ ràng buộc khóa chính, khóa ngoại, kiểm tra toàn vẹn và Indexing tối ưu tốc độ truy vấn.")
    add_bullet("ORM (Object-Relational Mapping): ", "Entity Framework Core 8.0 (Microsoft.EntityFrameworkCore.SqlServer). Hỗ trợ Code-First / DbContext với tính năng tự động khởi tạo bảng và tự nạp dữ liệu mẫu (Seed Data) nếu CSDL rỗng.")
    add_bullet("Thư viện xử lý Excel: ", "ClosedXML (v0.102.3) - Thư viện mã nguồn mở hàng đầu cho định dạng file .xlsx, cho phép tạo báo cáo có Header màu sắc chuyên nghiệp, viền nét mảnh, tự động căn độ rộng cột và kiểm tra lỗi từng ô dữ liệu.")

    add_h2("2.2. Mô hình kiến trúc 3 lớp (3-Layer Architecture)")
    add_p("Hệ thống tuân thủ chặt chẽ nguyên lý thiết kế phần mềm với kiến trúc 3 lớp phân tách rõ ràng trách nhiệm:")
    add_bullet("1. Tầng Giao diện (Presentation Layer - Forms): ", "Chứa 12 Windows Forms chuyên biệt và các Form hỗ trợ. Đảm nhận hiển thị dữ liệu, nhận sự kiện tương tác của người dùng, kiểm tra sơ bộ dữ liệu nhập (Client Validation) và chuyển giao yêu cầu cho tầng nghiệp vụ. Hoàn toàn không viết câu lệnh truy vấn SQL trực tiếp trong Form.")
    add_bullet("2. Tầng Nghiệp vụ (Business Logic Layer - Services & DTOs): ", "Gồm các dịch vụ: CustomerService, GameService, ParkServiceService, EmployeeService, TicketService, InvoiceService, ReportService, AuthService, ExcelService và ValidationHelper. Đây là nơi kiểm tra nghiệp vụ, kiểm tra ràng buộc số điện thoại 10 số, tính toán chiết khấu, trừ tồn kho dịch vụ trong Transaction và xuất nhập dữ liệu.")
    add_bullet("3. Tầng Truy xuất Dữ liệu (Data Access Layer - DAL): ", "Chứa QuanLyKhuVuiChoiDbContext, các Entity Models và DbInitializer. Đảm bảo kết nối an toàn với SQL Server, thực thi các truy vấn LINQ to Entities tối ưu, quản lý Transaction và ràng buộc toàn vẹn.")

    # =========================================================================
    # 3. THIẾT KẾ CƠ SỞ DỮ LIỆU CHI TIẾT
    # =========================================================================
    add_h1("3. THIẾT KẾ CƠ SỞ DỮ LIỆU CHI TIẾT")
    add_p("Cơ sở dữ liệu hệ thống được đặt tên là 'QuanLyKhuVuiChoi', gồm 8 bảng quan hệ liên kết chặt chẽ:")

    db_tables_summary = [
        ["1", "KhachHang", "MaKH (PK), HoTen, NgaySinh, GioiTinh, SoDienThoai (Unique), Email, DiaChi, NgayDangKy", "Lưu trữ hồ sơ khách hàng, số điện thoại là duy nhất để nhận diện."],
        ["2", "NhanVien", "MaNV (PK), HoTen, NgaySinh, GioiTinh, SoDienThoai, Email, DiaChi, ChucVu, NgayVaoLam, TrangThai", "Quản lý danh sách nhân sự, chức vụ và trạng thái làm việc."],
        ["3", "TaiKhoan", "TenDangNhap (PK), MatKhau, MaNV (FK), Quyen (Admin/NhanVien), TrangThai", "Lưu tài khoản đăng nhập, liên kết 1-1 với Nhân viên, phân quyền truy cập."],
        ["4", "TroChoi", "MaTroChoi (PK), TenTroChoi, LoaiTroChoi, GiaVe, DoTuoiToiThieu, SucChua, TrangThai, MoTa", "Danh mục các trò chơi cảm giác mạnh, thiếu nhi, sức chứa và giá vé."],
        ["5", "DichVu", "MaDichVu (PK), TenDichVu, LoaiDichVu, DonGia, SoLuongTon, TrangThai, MoTa", "Danh mục ẩm thực, nước uống, thuê đồ bơi, tủ gửi đồ; quản lý số lượng tồn."],
        ["6", "Ve", "MaVe (PK), TenVe, LoaiVe, GiaVe, ThoiHan, MoTa, TrangThai", "Bảng giá các loại vé vào cổng, vé trọn gói, vé VIP, vé ngày/vé tháng."],
        ["7", "HoaDon", "MaHoaDon (PK), MaKH (FK), MaNV (FK), NgayLap, TongTien, GiamGia, PhuongThucThanhToan, TrangThai, GhiChu", "Hóa đơn thanh toán bán hàng, lưu tổng tiền thực thu và phần trăm giảm giá."],
        ["8", "ChiTietHoaDon", "MaCTHD (PK Identity), MaHoaDon (FK), LoaiHang, MaHang, TenHang, SoLuong, DonGia, ThanhTien", "Chi tiết các mặt hàng (Vé, Dịch vụ, Trò chơi) thuộc về từng hóa đơn."]
    ]
    t_sum = doc.add_table(rows=9, cols=4)
    format_table(
        t_sum,
        [0.4, 1.2, 2.7, 2.4],
        ["STT", "Tên bảng", "Các trường thông tin chính", "Ý nghĩa nghiệp vụ"],
        db_tables_summary,
        align_cols={0: WD_ALIGN_PARAGRAPH.CENTER}
    )

    add_h2("3.1. Các chỉ mục (Indexes) và Toàn vẹn dữ liệu")
    add_bullet("Tối ưu hóa tìm kiếm: ", "Hệ thống xây dựng sẵn các chỉ mục truy vấn: IX_KhachHang_SoDienThoai (Unique), IX_KhachHang_HoTen, IX_NhanVien_HoTen, IX_TroChoi_TenTroChoi, IX_DichVu_TenDichVu, IX_Ve_TenVe, IX_HoaDon_NgayLap giúp tốc độ lọc hàng chục nghìn dòng chỉ mất vài mili-giây.")
    add_bullet("Toàn vẹn tham chiếu (Referential Integrity): ", "Khi xóa Khách hàng hoặc Nhân viên đã phát sinh hóa đơn, tầng nghiệp vụ chủ động chặn lại kèm thông báo tường minh ('Không thể xóa khách hàng X vì đang tồn tại N hóa đơn liên quan').")
    add_bullet("Giao dịch an toàn (Database Transactions): ", "Các hành động thanh toán hóa đơn và hủy hóa đơn được bao bọc trong Transaction. Nếu việc trừ kho hoặc ghi chi tiết gặp lỗi, toàn bộ thao tác sẽ Rollback ngay lập tức, ngăn ngừa hoàn toàn hiện tượng lệch số liệu kế toán.")

    # =========================================================================
    # 4. ĐÁP ỨNG ĐẦY ĐỦ 6 THAO TÁC CỐT LÕI
    # =========================================================================
    add_h1("4. ĐÁP ỨNG ĐẦY ĐỦ 6 THAO TÁC NGHIỆP VỤ CỐT LÕI")
    add_p("Tất cả các phân hệ quản lý chính trong phần mềm đều được thiết kế chuẩn mực, đáp ứng 100% 6 thao tác bắt buộc:")

    operations_data = [
        [
            "1. LƯU / THÊM (Create)",
            "• Nhập liệu trên các ô nhập (TextBox, ComboBox, NumericUpDown).\n"
            "• Hệ thống tự động sinh mã ID thông minh kế tiếp (KH021, TC016, DV016, NV011, VE011, HD031) mà không yêu cầu người dùng gõ tay.\n"
            "• Kiểm tra Validation: Bắt buộc điền các trường chính, số điện thoại đúng 10 số (bắt đầu bằng 0), email chuẩn RFC, giá tiền >= 0, sức chứa > 0.\n"
            "• Lưu vào SQL Server qua EF Core và nạp lại bảng hiển thị ngay lập tức."
        ],
        [
            "2. SỬA (Update)",
            "• Nhấp chọn một dòng bất kỳ trên DataGridView, dữ liệu tự động 'bắn' lên các ô nhập liệu tương ứng.\n"
            "• Cho phép chỉnh sửa thông tin (ngoại trừ Khóa chính được khóa ReadOnly để đảm bảo an toàn).\n"
            "• Kiểm tra trùng lặp số điện thoại với người khác trước khi cho phép lưu.\n"
            "• Hiển thị thông báo xác nhận thành công và làm mới lưới dữ liệu."
        ],
        [
            "3. XÓA (Delete)",
            "• Chọn dòng dữ liệu cần xóa trên bảng, nhấn nút 'Xóa'.\n"
            "• Hệ thống hiển thị hộp thoại xác nhận cảnh báo MessageBox (Yes/No) để tránh xóa nhầm.\n"
            "• Kiểm tra ràng buộc quan hệ: Nếu Khách hàng, Trò chơi, Dịch vụ hoặc Vé đã có trong hóa đơn thì chặn lại và thông báo lý do rõ ràng.\n"
            "• Xóa an toàn và thông báo thành công."
        ],
        [
            "4. TÌM KIẾM (Search)",
            "• Cung cấp ô nhập từ khóa tìm kiếm nhanh trên tất cả các màn hình.\n"
            "• Hỗ trợ 3 chế độ tìm kiếm linh hoạt: 'Chứa (Contains)', 'Bắt đầu bằng (StartsWith)', 'Chính xác (Exact Match)'.\n"
            "• Tìm kiếm đồng thời trên nhiều trường: Mã, Họ tên, Số điện thoại, Email hoặc Loại dịch vụ."
        ],
        [
            "5. XUẤT FILE EXCEL (Export)",
            "• Sử dụng thư viện ClosedXML xuất dữ liệu ra file định dạng Excel chuẩn (.xlsx).\n"
            "• Header báo cáo chuyên nghiệp: Banner màu xanh Navy (#1F4E78), chữ trắng in hoa cỡ 16pt, dòng hiển thị ngày giờ xuất và người thực hiện.\n"
            "• Định dạng tiền tệ đẹp mắt (#,##0 VNĐ), định dạng ngày chuẩn (dd/MM/yyyy).\n"
            "• Tự động căn chỉnh độ rộng cột vừa khít nội dung (AdjustToContents).\n"
            "• Hộp thoại hỏi người dùng có muốn mở file Excel vừa tạo ngay lập tức hay không."
        ],
        [
            "6. NHẬP FILE EXCEL (Import)",
            "• Pipeline xử lý nhập liệu chuyên nghiệp 5 bước:\n"
            "  + Bước 1: Chọn file .xlsx từ máy tính.\n"
            "  + Bước 2: ClosedXML đọc từng hàng và kiểm tra tính hợp lệ dữ liệu (Validate từng dòng).\n"
            "  + Bước 3: Phân loại: Dòng hợp lệ hiển thị lên bảng Preview xem trước; Dòng lỗi được liệt kê chi tiết (ví dụ: 'Dòng 5: SĐT không đúng 10 chữ số', 'Dòng 8: Trùng mã đã tồn tại').\n"
            "  + Bước 4: Người dùng bấm 'Xác nhận nhập' để lưu các dòng hợp lệ vào SQL Server.\n"
            "  + Bước 5: Thông báo số lượng bản ghi thêm mới thành công.\n"
            "• Cung cấp nút 'Tải file mẫu Excel' để người dùng xuất file mẫu chuẩn cấu trúc cột."
        ]
    ]

    t_ops = doc.add_table(rows=7, cols=2)
    format_table(
        t_ops,
        [2.0, 4.7],
        ["Thao tác cốt lõi", "Quy trình thực thi chi tiết"],
        operations_data
    )

    doc.add_page_break()

    # =========================================================================
    # 5. CHI TIẾT 12 FORMS BẮT BUỘC VÀ CÁC PHÂN HỆ
    # =========================================================================
    add_h1("5. ĐẶC TẢ CHI TIẾT CÁC PHÂN HỆ VÀ 12 MÀN HÌNH CHỨC NĂNG")
    add_p("Hệ thống được thiết kế theo đúng yêu cầu đề bài gồm 12 Forms chính phân chia đồng đều cho 4 thành viên trong nhóm, kết hợp các màn hình điều hướng và xác thực:")

    # 5.1. Auth & Shell
    add_h2("5.1. Phân hệ Xác thực & Điều hướng chung")
    
    add_h3("5.1.1. Form Đăng nhập hệ thống (FormDangNhap)")
    add_bullet("Mục đích: ", "Xác thực danh tính người dùng trước khi vào phần mềm, phân quyền làm việc theo vai trò Admin hoặc NhanVien.")
    add_bullet("Giao diện & Điều khiển: ", "Ô nhập Tên đăng nhập, Mật khẩu (ký tự dấu chấm bảo mật), nút 'Đăng nhập', nút 'Thoát', checkbox hiển thị mật khẩu.")
    add_bullet("Quy trình xử lý: ", "Kiểm tra tài khoản trong bảng TaiKhoan; so khớp mật khẩu và kiểm tra trạng thái kích hoạt. Nếu hợp lệ, lưu thông tin phiên làm việc vào AuthService (CurrentAccount, CurrentEmployee) và mở FormMain; nếu sai, hiển thị cảnh báo lỗi màu đỏ.")

    add_h3("5.1.2. Màn hình chính điều hướng (FormMain)")
    add_bullet("Mục đích: ", "Đóng vai trò là Master Layout (Giao diện điều khiển trung tâm) ứng dụng kiến trúc Form nhúng (Child Form Embedding).")
    add_bullet("Giao diện & Điều khiển: ", "Sidebar bên trái gồm Logo FunZone, danh mục Menu chuyển trang, Avatar nhân viên, tên người dùng đăng nhập và nút Đăng xuất. Topbar hiển thị lời chào theo tên nhân viên, đồng hồ thời gian thực (Timer cập nhật từng giây) và ô tìm kiếm nhanh. Vùng Content ở giữa nhúng mượt mà các Form chức năng.")
    add_bullet("Ưu điểm nổi bật: ", "Không mở cửa sổ dạng pop-up rời rạc gây rối mắt; người dùng chuyển qua lại giữa Bán vé, Khách hàng, Trò chơi mượt mà với thanh menu đổi màu Active trực quan.")

    add_h3("5.1.3. Bảng điều khiển tổng quan (FormDashboard)")
    add_bullet("Mục đích: ", "Cung cấp cho Quản trị viên và Quản lý ca cái nhìn toàn cảnh về tình hình kinh doanh trong ngày và trong tháng.")
    add_bullet("Giao diện & Điều khiển: ", "4 thẻ chỉ số KPI nổi bật (Thẻ Doanh thu hôm nay, Thẻ Doanh thu tháng này, Thẻ Số vé đã bán, Thẻ Tổng khách hàng thân thiết); Bảng DataGridView hiển thị danh sách 10 hóa đơn giao dịch mới nhất; Các nút lối tắt (Quick Actions) chuyển nhanh tới Bán vé, Thêm khách hàng, Xuất báo cáo.")

    # 5.2. THÀNH VIÊN 1
    add_h2("5.2. Phân hệ Quản lý Khách hàng (Thành viên 1 phụ trách)")

    add_h3("5.2.1. Form Quản lý Khách hàng (FormKhachHang - Form 1)")
    add_bullet("Mục đích: ", "Quản lý toàn bộ vòng đời thông tin khách hàng đến vui chơi tại công viên.")
    add_bullet("Giao diện & Điều khiển: ", "Panel nhập liệu bên trái (Mã KH, Họ tên, Ngày sinh, Giới tính, Số điện thoại, Email, Địa chỉ, Ngày đăng ký); Thanh công cụ (Thêm, Sửa, Xóa, Làm mới); Khu vực tìm kiếm theo từ khóa và ComboBox chọn chế độ tìm kiếm (Contains / StartsWith / Exact); Bảng DataGridView danh sách khách hàng.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Tự động sinh mã KH tiếp theo (dạng KH001, KH002...); "
               "Ràng buộc số điện thoại bắt buộc 10 chữ số và không được trùng lặp trong hệ thống; "
               "Ràng buộc định dạng Email hợp lệ; "
               "Chặn xóa nếu khách hàng đã có hóa đơn giao dịch nhằm bảo toàn chứng từ kế toán.")

    add_h3("5.2.2. Form Nhập Khách hàng từ Excel (FormImportKhachHang - Form 2)")
    add_bullet("Mục đích: ", "Hỗ trợ nạp nhanh hàng trăm khách hàng từ danh sách đối tác, đoàn thể, tour du lịch thông qua file Excel.")
    add_bullet("Giao diện & Điều khiển: ", "Nút 'Chọn file Excel (.xlsx)', ô hiển thị đường dẫn, nút 'Tải file mẫu', DataGridView xem trước dữ liệu hợp lệ (Preview), ListBox danh sách cảnh báo lỗi từng dòng, thanh thông báo trạng thái, nút 'Xác nhận nạp vào CSDL'.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "ClosedXML duyệt từng dòng từ hàng thứ 2; "
               "Kiểm tra tính hợp lệ: Bắt buộc có Họ tên và SĐT; SĐT đủ 10 số; kiểm tra không trùng mã hoặc trùng SĐT đã tồn tại trong CSDL hoặc trùng ngay trong chính file Excel; "
               "Hiển thị chi tiết vị trí lỗi trên ListBox để người dùng sửa file trước khi bấm xác nhận.")

    add_h3("5.2.3. Form Xuất Khách hàng ra Excel (FormExportKhachHang - Form 3)")
    add_bullet("Mục đích: ", "Xuất danh sách khách hàng ra file Excel phục vụ chiến dịch Marketing, gửi lời chúc sinh nhật hoặc lưu trữ báo cáo.")
    add_bullet("Giao diện & Điều khiển: ", "Bộ lọc dữ liệu (Xuất toàn bộ khách hàng, hoặc Xuất theo kết quả tìm kiếm/lọc), ô chọn đường dẫn lưu file (SaveFileDialog), nút 'Tiến hành xuất Excel'.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Tạo file Excel chuẩn bằng ClosedXML với Header Navy, Font Segoe UI; "
               "Tự động canh độ rộng cột; "
               "Mở hộp thoại hỏi: 'Xuất dữ liệu thành công! Bạn có muốn mở file Excel ngay bây giờ không?'.")

    # 5.3. THÀNH VIÊN 2
    add_h2("5.3. Phân hệ Quản lý Trò chơi & Dịch vụ (Thành viên 2 phụ trách)")

    add_h3("5.3.1. Form Quản lý Trò chơi (FormTroChoi - Form 4)")
    add_bullet("Mục đích: ", "Quản lý danh mục các trò chơi trong công viên, quy định an toàn và giá vé riêng lẻ từng trò chơi.")
    add_bullet("Giao diện & Điều khiển: ", "Các ô nhập liệu: Mã trò chơi, Tên trò chơi, Loại trò chơi (Cảm giác mạnh, Gia đình, Trẻ em, Dưới nước, Phiêu lưu), Giá vé lẻ, Độ tuổi tối thiểu, Sức chứa tối đa (người/lượt), Trạng thái (Đang hoạt động, Đang bảo trì, Tạm ngưng), Mô tả; Bảng DataGridView và bộ lọc theo loại trò chơi.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Tự động sinh mã trò chơi (TC001, TC002...); "
               "Ràng buộc giá vé >= 0, sức chứa >= 1, độ tuổi >= 0; "
               "Tìm kiếm theo tên trò chơi và loại trò chơi; "
               "Chặn xóa trò chơi nếu đã có chi tiết hóa đơn khách mua vé vào chơi.")

    add_h3("5.3.2. Form Quản lý Dịch vụ (FormDichVu - Form 5)")
    add_bullet("Mục đích: ", "Quản lý danh mục ẩm thực, nước uống, đồ lưu niệm, dịch vụ cho thuê đồ (tủ gửi đồ, chòi nghỉ, phao bơi) và số lượng hàng tồn kho.")
    add_bullet("Giao diện & Điều khiển: ", "Các trường: Mã dịch vụ, Tên dịch vụ, Phân loại (Ẩm thực, Nước uống, Thuê đồ, Khu nghỉ ngơi, Khác), Đơn giá, Số lượng tồn kho, Trạng thái (Đang phục vụ, Hết hàng, Tạm ngưng); Tính năng 'Cập nhật nhanh số lượng tồn' tiện lợi.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Tự động đổi trạng thái sang 'Hết hàng' khi số lượng tồn kho giảm về 0; "
               "Khi bán dịch vụ tại quầy POS, hệ thống tự động trừ kho; "
               "Ràng buộc đơn giá và tồn kho không được âm; "
               "Tự động sinh mã DV001, DV002...")

    add_h3("5.3.3. Form Nhập/Xuất Excel Trò chơi & Dịch vụ (FormImportExportTroChoiDichVu - Form 6)")
    add_bullet("Mục đích: ", "Cung cấp trung tâm tích hợp nhập/xuất file Excel cho cả 2 danh mục Trò chơi và Dịch vụ trong cùng một giao diện TabControl hiện đại.")
    add_bullet("Giao diện & Điều khiển: ", "Tab 1: Quản lý Import/Export Trò chơi; Tab 2: Quản lý Import/Export Dịch vụ. Mỗi Tab đều có đầy đủ: Nút Tải file mẫu, Xem trước (Preview) dòng hợp lệ, ListBox thông báo dòng lỗi, Nút Lưu vào CSDL và Nút Xuất Excel.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Validate giá vé, độ tuổi, sức chứa hợp lệ khi import trò chơi; "
               "Validate số lượng tồn kho và đơn giá khi import dịch vụ; "
               "Tải sẵn các file mẫu chuẩn trong thư mục ExcelTemplates.")

    # 5.4. THÀNH VIÊN 3
    add_h2("5.4. Phân hệ Quản lý Nhân sự & Bảng giá Vé (Thành viên 3 phụ trách)")

    add_h3("5.4.1. Form Quản lý Nhân viên (FormNhanVien - Form 7)")
    add_bullet("Mục đích: ", "Quản lý hồ sơ nhân viên trong công viên (Thu ngân, Soát vé, Kỹ thuật bảo trì, Chăm sóc khách hàng, Quản lý).")
    add_bullet("Giao diện & Điều khiển: ", "Các ô nhập: Mã NV, Họ tên, Ngày sinh, Giới tính, Số điện thoại, Email, Địa chỉ, Chức vụ, Ngày vào làm, Trạng thái làm việc (Đang làm việc, Đã nghỉ việc); Bộ lọc theo Chức vụ; Tìm kiếm đa tiêu chí.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Tự động sinh mã NV (NV001, NV002...); "
               "Kiểm tra số điện thoại đúng 10 số; "
               "Chặn xóa nhân viên nếu nhân viên đó đã đứng tên xuất hóa đơn bán hàng nhằm bảo toàn lịch sử thu ngân.")

    add_h3("5.4.2. Form Quản lý Bảng giá Vé (FormVe - Form 8)")
    add_bullet("Mục đích: ", "Thiết lập bảng giá các loại vé vào cổng và vé trọn gói của khu vui chơi.")
    add_bullet("Giao diện & Điều khiển: ", "Các trường: Mã vé, Tên vé, Loại vé (Vé ngày, Vé trẻ em, Vé người lớn, Vé gia đình, Vé VIP), Giá vé, Thời hạn hiệu lực (Trong ngày, 1 tháng...), Mô tả quyền lợi, Trạng thái (Đang áp dụng, Ngừng áp dụng).")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Tự động sinh mã VE001, VE002...; "
               "Giá vé bắt buộc >= 0; "
               "Cho phép kích hoạt / tạm ngừng loại vé để không hiển thị trên màn hình POS bán hàng khi hết chương trình ưu đãi.")

    add_h3("5.4.3. Form Nhập/Xuất Excel Nhân viên & Vé (FormImportExportNhanVienVe - Form 9)")
    add_bullet("Mục đích: ", "Công cụ nạp hàng loạt nhân sự theo đợt tuyển dụng mùa vụ và nạp bảng giá vé ưu đãi mới từ file Excel của phòng kinh doanh.")
    add_bullet("Giao diện & Điều khiển: ", "Thiết kế TabControl 2 phân vùng: Tab Nhân viên & Tab Loại vé. Đầy đủ tính năng Tải template, Duyệt file, Xem trước dữ liệu, Kiểm tra lỗi và Xuất báo cáo.")
    add_bullet("Nghiệp vụ chi tiết: ", 
               "Kiểm tra định dạng ngày vào làm, kiểm tra chức vụ hợp lệ; "
               "Kiểm tra giá vé không âm, tên vé không được để trống khi import.")

    # 5.5. THÀNH VIÊN 4
    add_h2("5.5. Phân hệ Bán vé POS, Hóa đơn & Báo cáo Tài chính (Thành viên 4 phụ trách)")

    add_h3("5.5.1. Form Tạo đơn Bán vé POS (FormBanVe - Form 10)")
    add_bullet("Mục đích: ", "Trọng tâm vận hành của công viên - Màn hình quầy bán vé (Point of Sale) tốc độ cao phục vụ khách mua vé và dịch vụ.")
    add_bullet("Giao diện & Điều khiển: ", 
               "Bên trái: Danh mục mặt hàng trực quan phân chia theo Tab (Tất cả, Vé vào cổng, Trò chơi, Dịch vụ ẩm thực). Mỗi mặt hàng hiển thị dưới dạng 'Thẻ Card' có tên, giá tiền, số lượng tồn kho và nút 'Thêm vào giỏ' màu sắc nổi bật.\n"
               "Bên phải: Giỏ hàng trực quan (Cart Items) hiển thị từng mặt hàng đã chọn, nút tăng/giảm (+/-) số lượng, nút xóa từng dòng, nút xóa sạch giỏ; Khu vực chọn khách hàng, nhập phần trăm giảm giá (0 - 100%), chọn hình thức thanh toán (Tiền mặt, Chuyển khoản, Thẻ), ô ghi chú; Nút 'THANH TOÁN & IN HÓA ĐƠN' lớn, rõ ràng.")
    add_bullet("Quy trình xử lý nghiệp vụ: ", 
               "• Tự động tính toán Tạm tính, Tiền giảm giá, và Tổng tiền cần thanh toán theo thời gian thực mỗi khi thay đổi số lượng mặt hàng hoặc % giảm giá.\n"
               "• Kiểm tra tồn kho trước khi cho thêm vào giỏ đối với Dịch vụ (không cho bán vượt quá số lượng còn lại trong kho).\n"
               "• Thực hiện lưu trong Transaction: Sinh mã hóa đơn mới (HDxxx) -> Tạo bản ghi HoaDon -> Tạo các bản ghi ChiTietHoaDon -> Tự động trừ tồn kho mặt hàng Dịch vụ -> Commit Transaction.\n"
               "• Tự động hiển thị cửa sổ FormChiTietHoaDon xem trước hóa đơn thanh toán để in cho khách.")

    add_h3("5.5.2. Form Quản lý Hóa đơn & Lịch sử giao dịch (FormHoaDon - Form 11)")
    add_bullet("Mục đích: ", "Quản lý tập trung toàn bộ các giao dịch đã thực hiện trong công viên, tra cứu và xử lý hủy đơn khi có khiếu nại.")
    add_bullet("Giao diện & Điều khiển: ", "Bộ lọc theo khoảng ngày (Từ ngày - Đến ngày), lọc theo Khách hàng, ô tìm kiếm mã hóa đơn hoặc tên thu ngân; DataGridView hiển thị danh sách hóa đơn; Các nút chức năng: 'Xem chi tiết hóa đơn', 'Hủy hóa đơn', 'Xuất Excel danh sách'.")
    add_bullet("Nghiệp vụ Hủy đơn an toàn (Cancel Invoice): ", 
               "Khi khách trả vé hoặc giao dịch bị hủy, người dùng bấm 'Hủy hóa đơn' -> Nhập lý do hủy đơn. "
               "Hệ thống thực hiện trong Transaction: Cập nhật trạng thái hóa đơn sang 'Đã hủy', ghi nhận lý do vào Ghi chú, và TỰ ĐỘNG CỘNG HOÀN TRẢ SỐ LƯỢNG TỒN KHO cho các mặt hàng dịch vụ thuộc hóa đơn đó. Bảo đảm kho hàng luôn khớp với thực tế.")

    add_h3("5.5.3. Form Chi tiết Hóa đơn (FormChiTietHoaDon - Màn hình dùng chung)")
    add_bullet("Mục đích: ", "Hiển thị chi tiết phiếu thanh toán hóa đơn đầy đủ như một hóa đơn điện tử / phiếu in nhiệt: Mã HĐ, Tên khách hàng, Nhân viên lập, Ngày giờ mua, Bảng chi tiết danh mục (Vé/Dịch vụ, Đơn giá, Số lượng, Thành tiền), Phần trăm giảm giá, Phương thức thanh toán và Tổng tiền cuối cùng. Cung cấp nút 'In hóa đơn / Đóng'.")

    add_h3("5.5.4. Form Báo cáo Doanh thu & Thống kê (FormBaoCaoDoanhThu - Form 12)")
    add_bullet("Mục đích: ", "Cung cấp báo cáo tài chính định kỳ giúp Ban Quản lý nắm bắt chính xác doanh số bán vé và dịch vụ.")
    add_bullet("Giao diện & Điều khiển: ", "Bộ lọc kỳ báo cáo linh hoạt (Hôm nay, Tuần này, Tháng này, Năm nay, Tùy chọn khoảng ngày); 3 thẻ tổng kết tài chính (Tổng doanh thu, Tổng số vé đã bán, Tổng số hóa đơn); 2 bảng dữ liệu phân tích: Bảng 1 thống kê doanh thu theo từng mốc thời gian; Bảng 2 phân tích cơ cấu tỷ trọng doanh thu theo từng nhóm (Loại Vé, Dịch vụ ăn uống, Vé trò chơi riêng); Nút 'Xuất báo cáo Doanh thu ra Excel'.")
    add_bullet("Nghiệp vụ xuất Excel tài chính: ", 
               "Tạo file báo cáo tài chính bằng ClosedXML, định dạng dòng tổng kết đậm (Total Row), định dạng số tiền tệ chuẩn VNĐ (#,##0), hỗ trợ mở file trực tiếp.")

    doc.add_page_break()

    # =========================================================================
    # 6. QUY TRÌNH NGHIỆP VỤ TIÊU BIỂU (WORKFLOWS)
    # =========================================================================
    add_h1("6. CÁC QUY TRÌNH NGHIỆP VỤ TIÊU BIỂU")

    add_h2("6.1. Quy trình Bán vé & Thanh toán tại Quầy POS")
    add_bullet("Bước 1: Tiếp nhận nhu cầu: ", "Thu ngân mở màn hình '🎟️ Bán vé (POS)'.")
    add_bullet("Bước 2: Chọn mặt hàng: ", "Nhấp chọn loại vé vào cổng (Vé người lớn, Vé trẻ em, Combo gia đình...) hoặc dịch vụ ăn uống. Mặt hàng được tự động thêm vào giỏ bên phải.")
    add_bullet("Bước 3: Điều chỉnh số lượng & Chiết khấu: ", "Bấm nút [+] hoặc [-] để tăng giảm số lượng. Nếu khách hàng có voucher hoặc ưu đãi thành viên, nhập tỷ lệ giảm giá (ví dụ: 10%). Tổng tiền được tự động tính lại ngay lập tức.")
    add_bullet("Bước 4: Nhập thông tin & Thanh toán: ", "Chọn thông tin khách hàng từ danh sách (hoặc để trống nếu khách vãng lai), chọn hình thức thanh toán (Tiền mặt / Chuyển khoản / Quẹt thẻ) và nhấn nút 'THANH TOÁN & IN HÓA ĐƠN'.")
    add_bullet("Bước 5: Hoàn tất & In phiếu: ", "Cơ sở dữ liệu lưu hóa đơn, tự động trừ tồn kho dịch vụ trong Database Transaction và bật màn hình xem chi tiết hóa đơn để in cho khách.")

    add_h2("6.2. Quy trình Nhập dữ liệu hàng loạt từ Excel (Import Pipeline)")
    add_bullet("Bước 1: Chuẩn bị file: ", "Người dùng nhấn 'Tải file mẫu Excel' để tải file mẫu chuẩn cấu trúc (có sẵn các cột tiêu đề quy định). Điền dữ liệu cần thêm vào file.")
    add_bullet("Bước 2: Chọn file nạp: ", "Nhấn 'Chọn file Excel (.xlsx)' và chọn file đã chuẩn bị.")
    add_bullet("Bước 3: Kiểm tra tự động (Validation): ", "Hệ thống duyệt từng dòng dữ liệu: Kiểm tra dòng rỗng, kiểm tra số điện thoại đúng 10 số, kiểm tra định dạng email, kiểm tra trùng khóa chính hoặc trùng SĐT.")
    add_bullet("Bước 4: Xem trước kết quả: ", "Dữ liệu hợp lệ xuất hiện trên DataGridView Preview; các dòng bị lỗi được thống kê rõ ràng tại khung thông báo ('Dòng X: Lỗi Y').")
    add_bullet("Bước 5: Xác nhận lưu: ", "Người dùng kiểm tra lại và nhấn nút 'Xác nhận nạp dữ liệu'. Hệ thống tiến hành ghi dữ liệu hợp lệ vào CSDL và thông báo tổng số dòng đã thêm thành công.")

    add_h2("6.3. Quy trình Xử lý Hủy hóa đơn & Hoàn trả kho an toàn")
    add_bullet("Bước 1: Tra cứu đơn: ", "Mở màn hình '💳 Thanh toán / Quản lý hóa đơn', tìm kiếm theo mã hóa đơn cần hủy.")
    add_bullet("Bước 2: Kiểm tra trạng thái: ", "Nếu hóa đơn đã bị hủy trước đó, hệ thống sẽ cảnh báo không cho hủy lại.")
    add_bullet("Bước 3: Nhập lý do & Xác nhận: ", "Nhấn nút 'Hủy hóa đơn', hộp thoại yêu cầu nhập lý do hủy đơn (ví dụ: 'Khách đổi lịch trình', 'Mua nhầm loại vé').")
    add_bullet("Bước 4: Thực thi giao dịch hoàn kho: ", "Hệ thống mở Database Transaction: Chuyển trạng thái sang 'Đã hủy', ghi chú lý do, và tự động cộng hoàn trả lại số lượng tồn kho cho toàn bộ mặt hàng dịch vụ ẩm thực thuộc đơn đó. Bảo đảm kho hàng không bị sai lệch.")

    # =========================================================================
    # 7. PHÂN CHIA TRÁCH NHIỆM CHO 4 THÀNH VIÊN
    # =========================================================================
    add_h1("7. BẢNG PHÂN CHIA CÔNG VIỆC CHO 4 THÀNH VIÊN")
    add_p("Toàn bộ 12 màn hình chức năng và các tầng dịch vụ liên quan được phân chia rõ ràng, minh bạch cho 4 thành viên trong nhóm:")

    team_data = [
        [
            "THÀNH VIÊN 1",
            "Quản lý Khách hàng",
            "1. FormKhachHang\n2. FormImportKhachHang\n3. FormExportKhachHang\n- CustomerService.cs",
            "• Thiết kế giao diện và xử lý CRUD hồ sơ khách hàng.\n• Tìm kiếm đa chế độ (Contains, StartsWith, Exact Match).\n• Nhập Excel kiểm tra SĐT 10 số, trùng mã, trùng SĐT, hiển thị lỗi từng dòng và xem trước (Preview).\n• Xuất Excel định dạng ClosedXML chuyên nghiệp."
        ],
        [
            "THÀNH VIÊN 2",
            "Quản lý Trò chơi & Dịch vụ",
            "4. FormTroChoi\n5. FormDichVu\n6. FormImportExportTroChoiDichVu\n- GameService.cs\n- ParkServiceService.cs",
            "• Quản lý danh mục trò chơi, độ tuổi, sức chứa, giá vé.\n• Quản lý dịch vụ ăn uống, tồn kho, tính năng cập nhật tồn nhanh, tự đổi trạng thái hết hàng.\n• Tích hợp trung tâm Nhập/Xuất Excel bằng TabControl cho Trò chơi và Dịch vụ kèm file mẫu."
        ],
        [
            "THÀNH VIÊN 3",
            "Quản lý Nhân viên & Bảng giá Vé",
            "7. FormNhanVien\n8. FormVe\n9. FormImportExportNhanVienVe\n- EmployeeService.cs\n- TicketService.cs",
            "• Quản lý hồ sơ nhân viên, chức vụ, trạng thái làm việc, kiểm tra số điện thoại.\n• Quản lý bảng giá các loại vé vào cổng, vé combo, thời hạn hiệu lực.\n• Xây dựng quy trình Nhập/Xuất Excel cho Nhân viên và Vé có kiểm tra toàn vẹn."
        ],
        [
            "THÀNH VIÊN 4",
            "Bán vé POS, Hóa đơn & Doanh thu",
            "10. FormBanVe (POS)\n11. FormHoaDon\n12. FormBaoCaoDoanhThu\n- Common: FormChiTietHoaDon\n- InvoiceService.cs\n- ReportService.cs",
            "• Màn hình bán vé POS hiện đại: Card sản phẩm, Giỏ hàng động, tính chiết khấu, phương thức thanh toán, in hóa đơn.\n• Quản lý hóa đơn: Tra cứu, xem chi tiết, hủy đơn an toàn có hoàn trả tồn kho trong Transaction.\n• Báo cáo doanh thu theo mốc thời gian, cơ cấu doanh thu theo nhóm hàng, xuất báo cáo ClosedXML."
        ]
    ]

    t_team = doc.add_table(rows=5, cols=4)
    format_table(
        t_team,
        [1.3, 1.4, 1.8, 2.2],
        ["Thành viên", "Phân hệ chính", "Các Form & Dịch vụ phụ trách", "Trách nhiệm và Tiêu chí hoàn thành"],
        team_data
    )

    doc.add_page_break()

    # =========================================================================
    # 8. HƯỚNG DẪN CÀI ĐẶT, VẬN HÀNH & TÀI KHOẢN MẪU
    # =========================================================================
    add_h1("8. HƯỚNG DẪN CÀI ĐẶT, VẬN HÀNH VÀ DỮ LIỆU DEMO")

    add_h2("8.1. Yêu cầu môi trường")
    add_bullet("Hệ điều hành: ", "Windows 10 / Windows 11 (64-bit).")
    add_bullet("Nền tảng .NET: ", ".NET 8.0 SDK trở lên (net8.0-windows).")
    add_bullet("Hệ quản trị CSDL: ", "Microsoft SQL Server 2017/2019/2022 hoặc SQL Server LocalDB / SQL Server Express.")
    add_bullet("Môi trường phát triển: ", "Visual Studio 2022 (cài đặt workload '.NET desktop development') hoặc VS Code kèm C# Dev Kit.")

    add_h2("8.2. Cấu hình chuỗi kết nối (Connection String)")
    add_p("Chuỗi kết nối CSDL được lưu trữ tập trung trong file appsettings.json tại thư mục gốc của dự án:")
    add_callout(
        '{\n'
        '  "ConnectionStrings": {\n'
        '    "DefaultConnection": "Server=localhost;Database=QuanLyKhuVuiChoi;Trusted_Connection=True;TrustServerCertificate=True;"\n'
        '  }\n'
        '}',
        title="Cấu hình chuỗi kết nối trong appsettings.json",
        alert_type="note"
    )
    add_p("Lưu ý: Nếu sử dụng SQL Server Express, người dùng chỉ cần thay đổi server thành 'Server=.\\\\SQLEXPRESS;Database=QuanLyKhuVuiChoi;...'.")

    add_h2("8.3. Cơ chế Khởi tạo CSDL và Tự nạp Dữ liệu mẫu (Auto Seeding)")
    add_p("Dự án được tích hợp cơ chế tự động hóa thông minh bậc nhất:")
    add_bullet("Tự động tạo bảng & Nạp mẫu: ", "Khi người dùng chạy chương trình lần đầu, nếu CSDL SQL Server chưa có dữ liệu, lớp DbInitializer sẽ tự động thực thi, tạo đầy đủ 8 bảng và nạp sẵn bộ dữ liệu mẫu vô cùng thực tế:")
    add_bullet("Dữ liệu mẫu phong phú: ", "20 Khách hàng đầy đủ thông tin địa chỉ tại Hà Nội; 10 Nhân viên các phòng ban; 15 Trò chơi mạo hiểm và thiếu nhi; 15 Dịch vụ ăn uống và tiện ích; 10 Loại vé vào cổng; 30 Hóa đơn giao dịch có sẵn đầy đủ chi tiết.")
    add_bullet("File SQL đính kèm: ", "Nếu muốn chủ động tạo CSDL qua SSMS, người dùng có thể mở và thực thi trực tiếp file Database/Database_Setup.sql có sẵn trong thư mục dự án.")

    add_h2("8.4. Danh sách tài khoản đăng nhập thử nghiệm")
    acc_data = [
        ["admin", "123456", "Admin (Toàn quyền)", "Trương Gia Bình", "Quản lý toàn quyền hệ thống, xem dashboard, chỉnh sửa mọi danh mục."],
        ["thungan1", "123456", "NhanVien", "Lê Thị Hương", "Thu ngân quầy POS bán vé, xem danh sách hóa đơn trong ca."],
        ["thungan2", "123456", "NhanVien", "Nguyễn Quang Hải", "Thu ngân quầy POS bán vé, phục vụ khách hàng."],
        ["soatve", "123456", "NhanVien", "Phạm Hải Yến", "Nhân viên soát vé và kiểm tra thông tin vé vào cổng."]
    ]
    t_acc = doc.add_table(rows=5, cols=5)
    format_table(
        t_acc,
        [1.0, 1.0, 1.3, 1.4, 2.0],
        ["Tài khoản", "Mật khẩu", "Phân quyền", "Nhân viên liên kết", "Ghi chú phân quyền"],
        acc_data,
        align_cols={0: WD_ALIGN_PARAGRAPH.CENTER, 1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER}
    )

    add_h2("8.5. Cách thức khởi chạy chương trình")
    add_p("Người dùng có thể khởi chạy phần mềm thông qua 3 cách cực kỳ đơn giản:")
    add_bullet("Cách 1 (Nhanh nhất): ", "Nhấp đúp chuột vào file CHAY_PHAN_MEM.bat có sẵn trong thư mục dự án. Script batch sẽ tự động gọi lệnh dotnet run và mở ứng dụng.")
    add_bullet("Cách 2 (Visual Studio 2022): ", "Mở thư mục dự án bằng Visual Studio 2022 và nhấn phím F5 hoặc nút Start.")
    add_bullet("Cách 3 (Command Line): ", "Mở PowerShell tại thư mục dự án và gõ lệnh: dotnet run.")

    # Save to file
    out_dir_1 = r"c:\Users\KIM THIEN\OneDrive\Desktop\Playground\QuanLyKhuVuiChoi"
    out_file_1 = os.path.join(out_dir_1, "TaiLieu_MoTa_Chi_Tiet_Chuc_Nang.docx")
    doc.save(out_file_1)
    print(f"Saved successfully to: {out_file_1}")

    out_dir_2 = r"c:\Users\KIM THIEN\OneDrive\Desktop\Playground"
    out_file_2 = os.path.join(out_dir_2, "TaiLieu_MoTa_Chi_Tiet_Chuc_Nang.docx")
    doc.save(out_file_2)
    print(f"Saved successfully to: {out_file_2}")

if __name__ == "__main__":
    create_document()
